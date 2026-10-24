"""Regression coverage for import integrity and safe project publication."""

import tempfile
import unittest
from pathlib import Path
from unittest import mock

from app.chapters.detector import Chapter, detect_chapters
from app.core.project import Project, create_project, list_projects
from app.documents.parser import (
    extract_text, clean_text, clean_import_noise, pdf_extraction_quality_warning,
    detect_repeated_book_noise, chapter_text_coverage, validate_chapter_text_coverage,
)


class ImportIntegrityTests(unittest.TestCase):
    def test_txt_import_strips_bom_and_preserves_story_content(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "story.txt"
            source.write_bytes(
                b"\xef\xbb\xbfChapter 1: Arrival\n"
                b"The hero walked toward the door.\n\n"
                b"Chapter 2: Departure\nShe followed him."
            )

            book = extract_text(source)
            chapters = detect_chapters(book.text)

            self.assertTrue(book.text.startswith("Chapter 1: Arrival"))
            self.assertIn("The hero walked toward the door.", book.text)
            self.assertIn("She followed him.", book.text)
            self.assertEqual([chapter.title for chapter in chapters], [
                "Chapter 1: Arrival", "Chapter 2: Departure"
            ])
            self.assertIn("The hero walked toward the door.", chapters[0].text)
            self.assertIn("She followed him.", chapters[1].text)

    def test_docx_import_preserves_paragraphs_and_table_story_text(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "story.docx"
            document = Document()
            document.add_paragraph("Chapter 1: Arrival")
            document.add_paragraph("The first paragraph survives.")
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Story text stored in a table survives."
            document.add_paragraph("The final paragraph survives.")
            document.save(source)

            book = extract_text(source)

            self.assertIn("The first paragraph survives.", book.text)
            self.assertIn("Story text stored in a table survives.", book.text)
            self.assertIn("The final paragraph survives.", book.text)
            self.assertLess(book.text.index("The first paragraph survives."),
                            book.text.index("Story text stored in a table survives."))
            self.assertLess(book.text.index("Story text stored in a table survives."),
                            book.text.index("The final paragraph survives."))

    def test_epub_import_follows_spine_reading_order(self):
        from ebooklib import epub

        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "story.epub"
            book = epub.EpubBook()
            book.set_identifier("import-integrity-test")
            book.set_title("Reading order test")
            book.set_language("en")
            chapter_two = epub.EpubHtml(title="Second", file_name="second.xhtml", lang="en")
            chapter_two.content = "<html><body><p>Second chapter content.</p></body></html>"
            chapter_one = epub.EpubHtml(title="First", file_name="first.xhtml", lang="en")
            chapter_one.content = "<html><body><p>First chapter content.</p></body></html>"
            book.add_item(chapter_one)
            book.add_item(chapter_two)
            book.spine = [chapter_two, chapter_one]
            book.add_item(epub.EpubNcx())
            book.add_item(epub.EpubNav())
            epub.write_epub(str(source), book)

            extracted = extract_text(source).text

            self.assertLess(extracted.index("Second chapter content."),
                            extracted.index("First chapter content."))

    def test_sparse_long_pdf_extraction_is_blocked_with_actionable_diagnostic(self):
        warning = pdf_extraction_quality_warning(page_count=80, word_count=120)
        self.assertIsNotNone(warning)
        self.assertIn("OCR", warning)
        self.assertIn("partial audiobook", warning)
        self.assertIsNone(pdf_extraction_quality_warning(page_count=2, word_count=120))

    def test_chapter_detection_does_not_leak_next_heading_into_previous_body(self):
        source = (
            "Chapter 1: First\n"
            "First chapter text ends here.\n\n"
            "Chapter 2: Second\n"
            "Second chapter text ends here."
        )
        chapters = detect_chapters(source)

        self.assertEqual(len(chapters), 2)
        self.assertNotIn("Chapter 2: Second", chapters[0].text)
        self.assertNotIn("Second chapter text", chapters[0].text)
        self.assertIn("First chapter text ends here.", chapters[0].text)
        self.assertIn("Second chapter text ends here.", chapters[1].text)

    def test_realistic_multi_chapter_watermark_cleanup_keeps_story_and_removes_page_noise(self):
        chapters_text = [
            "Chapter 1: The Arrival\nA heroine followed a long road toward the old city gate.\n41\nReport\nwww.asianovel.com",
            "Chapter 2: The Return\nThe hero came back from the northern pass after the battle ended.\n42\nReport\nwww.asianovel.com",
            "Chapter 3: The Promise\nThey promised to meet again before the next winter festival.\n43\nReport\nwww.asianovel.com",
        ]
        repeated = set(detect_repeated_book_noise(chapters_text))
        self.assertIn("report", repeated)
        self.assertIn("www.asianovel.com", repeated)

        cleaned_chapters = []
        removed_all = []
        for text in chapters_text:
            cleaned, removed = clean_import_noise(text, repeated)
            cleaned_chapters.append(cleaned)
            removed_all.extend(removed)

        story = "\n".join(cleaned_chapters)
        for expected in ("heroine followed", "hero came back", "promised to meet"):
            self.assertIn(expected, story)
        for noise in ("www.asianovel.com", "Report", "41", "42", "43"):
            self.assertNotIn(noise, story)
        self.assertGreaterEqual(len(removed_all), 9)

    def test_justlightnovels_page_header_is_removed_on_import_and_clear_text(self):
        source = (
            "Chapter 1: Arrival\n"
            "The traveler walked past the ancient gate and greeted the guard.\n"
            "Page | 42 www.justlightnovels.com\n"
            "She continued toward the village before sunset."
        )
        cleaned, removed = clean_import_noise(source)
        self.assertNotIn("justlightnovels.com", cleaned)
        self.assertNotIn("Page | 42", cleaned)
        self.assertTrue(any("justlightnovels.com" in item for item in removed))
        self.assertIn("greeted the guard", cleaned)
        self.assertIn("before sunset", cleaned)

        extracted = clean_text(source)
        self.assertNotIn("justlightnovels.com", extracted)
        self.assertIn("greeted the guard", extracted)
        self.assertIn("before sunset", extracted)

    def test_split_slash_title_and_part_line_are_combined_into_one_chapter_title(self):
        source = (
            "BLACK DOG/BIRTH\n\n"
            "Part 1\n"
            "The hero opened his eyes and wondered why the quiet room smelled of rain. "
            "He reached for the window and listened to the distant voices outside.\n\n"
            "Chapter 2: The Second Morning\n"
            "The next morning began with a clear sky and a long walk through town."
        )
        chapters = detect_chapters(source)
        self.assertEqual(chapters[0].title, "BLACK DOG/BIRTH: Part 1")
        self.assertTrue(chapters[0].text.startswith("BLACK DOG/BIRTH: Part 1\n"))
        self.assertIn("opened his eyes", chapters[0].text)
        self.assertNotIn("\nPart 1\n", chapters[0].text)
        self.assertTrue(any(chapter.title == "Chapter 2: The Second Morning" for chapter in chapters))

    def test_substantive_unheaded_story_opening_becomes_an_opening_chapter(self):
        opening = (
            "When he finally opened his eyes, the old station was empty and rain tapped "
            "against the cracked windows. He remembered the promise he had made yesterday, "
            "but the voice from the platform had already vanished."
        )
        source = opening + "\n\nChapter 1: A New Day\nThe sun rose over the town."
        chapters = detect_chapters(source)
        self.assertEqual(chapters[0].title, "Opening")
        self.assertIn("old station was empty", chapters[0].text)
        self.assertEqual(chapters[1].title, "Chapter 1: A New Day")
        self.assertNotIn("old station was empty", chapters[1].text)

    def test_single_occurrence_report_page_number_and_general_url_are_not_removed(self):
        source = (
            "Chapter 1: Archive\n"
            "The report was placed beside the old ledger.\n\n"
            "Report\n\n"
            "The portal at https://example.com/archive remained open in the story.\n\n"
            "17\n\n"
            "The keeper locked the library before sunrise."
        )
        cleaned, removed = clean_import_noise(source)
        self.assertIn("Report", cleaned)
        self.assertIn("17", cleaned)
        self.assertIn("https://example.com/archive", cleaned)
        self.assertEqual(removed, [])

    def test_chapter_coverage_guard_handles_short_book_and_detected_content_loss(self):
        source = (
            "Chapter 1: Begin\nThe first body keeps enough words to verify the small book.\n"
            "Chapter 2: End\nThe last body remains in the story and should be preserved."
        )
        chapters = detect_chapters(source)
        source_words, chapter_words, coverage = chapter_text_coverage(source, chapters)
        self.assertGreater(source_words, 0)
        self.assertGreater(chapter_words, 0)
        self.assertGreater(coverage, 0.90)
        validate_chapter_text_coverage(source, chapters)

        with self.assertRaisesRegex(ValueError, "Import was stopped"):
            validate_chapter_text_coverage(source, [Chapter(1, "Only", "one word")])

    def test_many_short_chapters_pass_coverage_guard_without_false_loss(self):
        parts = []
        for number in range(1, 13):
            parts.append(f"Chapter {number}: The Quiet Return")
            parts.append("The dawn breaks now.")
        source = "\n".join(parts)
        chapters = detect_chapters(source)

        self.assertEqual(len(chapters), 12)
        source_words, chapter_words, coverage = chapter_text_coverage(source, chapters)
        self.assertEqual(chapter_words, source_words)
        self.assertAlmostEqual(coverage, 1.0)
        validate_chapter_text_coverage(source, chapters, minimum=0.90)

    def test_txt_import_to_chapter_cleanup_pipeline_preserves_story_bodies(self):
        source_text = "\n\n".join([
            "Chapter 1: The Arrival\nA heroine followed the old road to the city gate.\n41\nReport\nwww.asianovel.com",
            "Chapter 2: The Return\nThe hero came home after the long battle ended.\n42\nReport\nwww.asianovel.com",
            "Chapter 3: The Promise\nThey promised to meet again before the winter festival.\n43\nReport\nwww.asianovel.com",
        ])
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "realistic-light-novel.txt"
            source.write_text(source_text, encoding="utf-8")
            book = extract_text(source)

        chapters = detect_chapters(book.text)
        self.assertEqual(len(chapters), 3)
        validate_chapter_text_coverage(book.text, chapters, minimum=0.90)
        repeated = set(detect_repeated_book_noise([chapter.text for chapter in chapters]))
        cleaned_chapters = []
        for chapter in chapters:
            chapter.text, _removed = clean_import_noise(chapter.text, repeated)
            cleaned_chapters.append(chapter)

        joined = "\n".join(chapter.text for chapter in cleaned_chapters)
        for body in ("heroine followed", "hero came home", "promised to meet"):
            self.assertIn(body, joined)
        self.assertNotIn("Report", joined)
        self.assertNotIn("41", joined)
        self.assertNotIn("42", joined)
        self.assertNotIn("43", joined)
        self.assertNotIn("www.asianovel.com", joined)

    def test_duplicate_dotted_toc_is_removed_and_passes_coverage_guard(self):
        source = """TABLE OF CONTENTS
Prologue ........ 1
Chapter 1: Arrival ........ 2
Chapter 2: Return ........ 3
Epilogue ........ 4

Prologue
The heroine walked through the old stone gate and greeted the guard before sunrise.
Chapter 1: Arrival
After the long journey, the traveler returned home and shared news with the waiting family.
Chapter 2: Return
They promised to meet again next winter, when the city would celebrate its annual winter festival.
Epilogue
In the final scene, the family gathered beside the river and remembered how their journey had changed everything.
"""
        chapters = detect_chapters(source)
        self.assertEqual(
            [chapter.title for chapter in chapters],
            ["Prologue", "Chapter 1: Arrival", "Chapter 2: Return", "Epilogue"],
        )
        joined = "\n".join(chapter.text for chapter in chapters)
        self.assertNotIn("TABLE OF CONTENTS", joined)
        self.assertNotIn("........", joined)
        source_words, chapter_words, coverage = chapter_text_coverage(source, chapters)
        self.assertEqual(source_words, chapter_words)
        self.assertAlmostEqual(coverage, 1.0)
        validate_chapter_text_coverage(source, chapters, minimum=0.90)

    def test_premium_v2_import_worker_persists_dotted_toc_book_end_to_end(self):
        from app.ui.premium_v2_host import PremiumV2Bridge

        source_text = """TABLE OF CONTENTS
Prologue ........ 1
Chapter 1: Arrival ........ 2
Chapter 2: Return ........ 3
Epilogue ........ 4

Prologue
The heroine walked through the old stone gate and greeted the guard before sunrise.
41
Report
www.asianovel.com

Chapter 1: Arrival
After the long journey, the traveler returned home and shared news with the waiting family.
42
Report
www.asianovel.com

Chapter 2: Return
They promised to meet again next winter, when the city would celebrate its annual winter festival.
43
Report
www.asianovel.com

Epilogue
In the final scene, the family gathered beside the river and remembered how their journey had changed everything.
44
Report
www.asianovel.com
"""
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "dotted-toc-book.txt"
            source.write_text(source_text, encoding="utf-8")
            app_data = root / "app-data"
            settings = root / "settings"
            library = root / "library"

            with (
                mock.patch("app.ui.premium_v2_host.ensure_roots"),
                mock.patch("app.ui.premium_v2_host.app_root", return_value=app_data),
                mock.patch("app.ui.premium_v2_host.settings_root", return_value=settings),
                mock.patch("app.ui.premium_v2_host.QSettings"),
                mock.patch("app.ui.premium_v2_host.load_state", return_value={}),
                mock.patch("app.ui.premium_v2_host.save_state"),
                mock.patch("app.core.project.library_root", return_value=library),
            ):
                bridge = PremiumV2Bridge()
                job_id = bridge._new_job("import", "Import dotted TOC book", "Acceptance test")
                bridge._job_update(
                    job_id, state="running", started_at=bridge._now(),
                    stage="Reading source file", progress=5,
                )
                bridge._import_worker(job_id, source, "Dotted TOC Integration Book")
                job = bridge._jobs[job_id]

                self.assertEqual(
                    job["state"], "completed",
                    f"Premium V2 import failed at {job.get('stage')}: {job.get('error') or job.get('message')}",
                )
                self.assertEqual(job["result"]["chapter_count"], 4)
                self.assertEqual(job["result"]["removed_noise_count"] > 0, True)

                project = Project.load(Path(job["project_folder"]))
                self.assertEqual(
                    [chapter.title for chapter in project.chapters],
                    ["Prologue", "Chapter 1: Arrival", "Chapter 2: Return", "Epilogue"],
                )
                joined = "\n".join(chapter.text for chapter in project.chapters)
                for expected in (
                    "heroine walked through the old stone gate",
                    "traveler returned home",
                    "promised to meet again next winter",
                    "family gathered beside the river",
                ):
                    self.assertIn(expected, joined)
                for noise in ("TABLE OF CONTENTS", "........", "www.asianovel.com", "Report", "41", "42", "43", "44"):
                    self.assertNotIn(noise, joined)

    def test_noise_cleanup_removes_known_watermark_without_dropping_story(self):
        source = (
            "The story mentions a report in the ordinary narration.\n"
            "41\nReport\nwww.asianovel.com\n"
            "The next paragraph remains intact."
        )
        cleaned, removed = clean_import_noise(source)

        self.assertNotIn("www.asianovel.com", cleaned)
        self.assertIn("The story mentions a report in the ordinary narration.", cleaned)
        self.assertIn("The next paragraph remains intact.", cleaned)
        self.assertIn("Report", removed)

    def test_standalone_report_word_is_preserved_without_watermark_context(self):
        source = "She handed him the report.\n\nReport\n\nThe meeting continued."
        cleaned, removed = clean_import_noise(source)

        self.assertIn("Report", cleaned)
        self.assertIn("The meeting continued.", cleaned)
        self.assertNotIn("Report", removed)

    def test_import_worker_applies_noise_cleanup_and_reports_removed_count(self):
        source = Path("app/ui/premium_v2_host.py").read_text(encoding="utf-8")
        self.assertIn("validate_chapter_text_coverage(text, chapters, minimum=0.90)", source)
        self.assertIn("detect_repeated_book_noise([chapter.text for chapter in chapters])", source)
        self.assertIn("chapter.text, removed = clean_import_noise(chapter.text, repeated_noise)", source)
        self.assertIn('"removed_noise_count": len(removed_noise)', source)

    def test_cleanup_preserves_urls_unicode_punctuation_and_dialogue_not_matching_watermarks(self):
        from app.documents.parser import clean_import_noise

        source = (
            "The character opened https://example.com/path to read the letter.\n"
            "“I said, report it!” she shouted.\n"
            "世界はまだ終わっていない。\n"
            "The hero kept walking—slowly, carefully.\n"
            "The next paragraph stays separate."
        )
        cleaned, removed = clean_import_noise(source)

        self.assertIn("https://example.com/path", cleaned)
        self.assertIn("“I said, report it!” she shouted.", cleaned)
        self.assertIn("世界はまだ終わっていない。", cleaned)
        self.assertIn("walking—slowly, carefully.", cleaned)
        self.assertIn("The next paragraph stays separate.", cleaned)
        self.assertEqual(removed, [])

    def test_repeated_generic_domain_is_preserved_as_story_content(self):
        source = [
            "Chapter 1: Archive\nThe character reads a note linking to www.example.com before leaving.",
            "Chapter 2: Portal\nA second character mentions www.example.com while describing the old city.",
        ]
        repeated = detect_repeated_book_noise(source)
        self.assertNotIn("www.example.com", repeated)
        cleaned, removed = clean_import_noise(source[0], set(repeated))
        self.assertIn("www.example.com", cleaned)
        self.assertEqual(removed, [])

    def test_reflow_joins_hyphenated_line_wrap_without_losing_paragraph_or_heading_boundaries(self):
        from app.documents.parser import reflow_source_text

        source = (
            "Chapter 1: Arrival\n"
            "The inter-\n"
            "national train slowed near the station.\n\n"
            "“Stay here,” she said.\n"
            "“I will.”\n\n"
            "Chapter 2: Departure\n"
            "The next part begins."
        )
        result = reflow_source_text(source)

        self.assertIn("Chapter 1: Arrival\nThe international train slowed near the station.", result)
        self.assertIn("“Stay here,” she said.\n“I will.”", result)
        self.assertIn("Chapter 2: Departure\nThe next part begins.", result)
        self.assertNotIn("Arrival The", result)
        self.assertNotIn("Chapter 1: Arrival Chapter 2: Departure", result)

    def test_chapter_detection_supports_roman_numerals_prologue_and_epilogue(self):
        source = (
            "Prologue\nA beginning before the main story.\n"
            "Chapter IV: The Trial\nThe trial starts.\n"
            "Epilogue\nThe story closes."
        )
        chapters = detect_chapters(source)

        self.assertEqual([chapter.title for chapter in chapters],
                         ["Prologue", "Chapter IV: The Trial", "Epilogue"])
        self.assertIn("A beginning before the main story.", chapters[0].text)
        self.assertIn("The trial starts.", chapters[1].text)
        self.assertIn("The story closes.", chapters[2].text)

    def test_chapter_detection_preserves_preamble_and_books_without_headings(self):
        preamble_source = (
            "Translator note: names follow the original spelling.\n"
            "Chapter 1: First\nThe first chapter begins."
        )
        chapters = detect_chapters(preamble_source)
        self.assertEqual(len(chapters), 1)
        self.assertIn("Translator note: names follow the original spelling.", chapters[0].text)
        self.assertIn("The first chapter begins.", chapters[0].text)

        unheaded = detect_chapters("A complete opening.\n\nA complete ending.")
        self.assertEqual(len(unheaded), 1)
        self.assertEqual(unheaded[0].title, "Full Book")
        self.assertIn("A complete opening.", unheaded[0].text)
        self.assertIn("A complete ending.", unheaded[0].text)

    def test_chapter_detection_does_not_treat_dialogue_like_chapter_text_as_heading(self):
        source = (
            "Chapter 1: Arrival\n"
            "Chapter 2: You cannot go in there, she said, turning away.\n"
            "She reached for the door."
        )
        chapters = detect_chapters(source)
        self.assertEqual(len(chapters), 1)
        self.assertEqual(chapters[0].title, "Chapter 1: Arrival")
        self.assertIn("Chapter 2: You cannot go in there, she said, turning away.", chapters[0].text)
        self.assertIn("She reached for the door.", chapters[0].text)

    def test_project_creation_round_trip_preserves_review_and_skip_state(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp) / "library"
            source = Path(temp) / "source.txt"
            source.write_text("A complete source book.", encoding="utf-8")
            chapter = Chapter(
                number=1,
                title="Chapter 1",
                text="A complete chapter.",
                review_status="reviewed",
                skip_for_generation=True,
            )

            with mock.patch("app.core.project.library_root", return_value=root):
                project = create_project("Test Book", source, [chapter])
                loaded = Project.load(project.folder)

            self.assertEqual(loaded.title, "Test Book")
            self.assertTrue(loaded.source_path.is_file())
            self.assertEqual(loaded.chapters[0].review_status, "reviewed")
            self.assertTrue(loaded.chapters[0].skip_for_generation)
            self.assertEqual(loaded.chapters[0].text, "A complete chapter.")

    def test_failed_project_metadata_or_chapter_write_removes_staging_folder(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp) / "library"
            source = Path(temp) / "source.txt"
            source.write_text("A complete source book.", encoding="utf-8")

            with mock.patch("app.core.project.library_root", return_value=root):
                with mock.patch.object(Project, "save", side_effect=OSError("simulated metadata write failure")):
                    with self.assertRaisesRegex(OSError, "simulated metadata write failure"):
                        create_project("Write Failure", source, [Chapter(1, "Chapter 1", "Text")])

            self.assertTrue(root.exists())
            self.assertEqual(list(root.iterdir()), [])
            self.assertEqual(list_projects(), [])

    def test_library_listing_hides_in_progress_import_staging_directories(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp) / "library"
            staging = root / ".import-staging-test"
            staging.mkdir(parents=True)
            (staging / "project.json").write_text('{"title":"Not published"}', encoding="utf-8")

            with mock.patch("app.core.project.library_root", return_value=root):
                self.assertEqual(list_projects(), [])

    def test_failed_source_copy_leaves_no_partial_project_in_library(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp) / "library"
            source = Path(temp) / "source.txt"
            source.write_text("A complete source book.", encoding="utf-8")

            with mock.patch("app.core.project.library_root", return_value=root):
                with mock.patch("app.core.project.shutil.copy2", side_effect=OSError("simulated copy failure")):
                    with self.assertRaisesRegex(OSError, "simulated copy failure"):
                        create_project("Broken Import", source, [Chapter(1, "Chapter 1", "Text")])

            self.assertTrue(root.exists())
            self.assertEqual(list(root.iterdir()), [])


if __name__ == "__main__":
    unittest.main()
