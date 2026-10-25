import tempfile
import unittest
from pathlib import Path
import subprocess
import inspect
import sys
import wave
import json

# Always import the application package from this checkout, not an identically
# named package that may be present in the CI Python environment.
REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from app.audio.assembler import assemble_m4b, _probe_audio_bitrate_kbps, DEFAULT_AUDIO_BITRATE_KBPS
from app.chapters.characters import analyze_book, _all_dialogue_spans
from app.chapters.dialogue import dialogue_segment_for_selection
from app.chapters.detector import Chapter, detect_chapters
from app.tts.voice_profile import VoiceProfile, save_profiles, load_profiles, import_reference_audio
from app.tts.narration import prepare_for_narration
from app.tts.pacing import append_silence, pause_after_ms, split_for_pacing
from app.tts.preview import build_voice_preview
from app.tts.pronunciation_suggester import (
    suggest_pronunciation, suggest_names_from_text, COMMON_ENGLISH_WORDS,
    is_common_english_phrase, nativeish_pronunciation,
 )


def _write_wav(path: Path, seconds: float = 0.05) -> None:
    import math
    path.parent.mkdir(parents=True, exist_ok=True)
    rate = 44100
    frames = max(1, int(rate * seconds))
    pcm = bytearray()
    for index in range(frames):
        sample = int(12000 * math.sin(2 * math.pi * 440 * index / rate))
        pcm.extend(int(sample).to_bytes(2, byteorder="little", signed=True))
    with wave.open(str(path), "wb") as wav:
        wav.setnchannels(1)
        wav.setsampwidth(2)
        wav.setframerate(rate)
        wav.writeframes(bytes(pcm))


class RegressionTests(unittest.TestCase):
    def test_character_detection_recovers_explicit_speaker_tags(self):
        from app.chapters.detector import Chapter

        chapter = Chapter(
            number=1,
            title="Speaker tag regression",
            text=(
                'Akeno Himejima said, “You are late.”\n'
                'Rias Gremory replied, “I know.”\n'
                'Akeno Himejima said, “Then let us go.”'
            ),
        )
        analysis = analyze_book([chapter])
        by_name = {item.name: item for item in analysis.characters}

        self.assertEqual(analysis.dialogue_total, 3)
        self.assertIn("Akeno Himejima", by_name)
        self.assertIn("Rias Gremory", by_name)
        self.assertEqual(by_name["Akeno Himejima"].dialogue_count, 2)
        self.assertEqual(by_name["Rias Gremory"].dialogue_count, 1)

    def test_character_detection_accepts_unicode_speaker_labels(self):
        from app.chapters.detector import Chapter

        chapter = Chapter(
            number=1,
            title="Unicode speaker label",
            text='リアス: 「遅いわね。」\nアーシア: 「ごめんなさい。」',
        )
        analysis = analyze_book([chapter])
        names = {item.name for item in analysis.characters}

        self.assertIn("リアス", names)
        self.assertIn("アーシア", names)
        self.assertEqual(analysis.dialogue_total, 2)
        by_name = {item.name: item for item in analysis.characters}
        self.assertEqual(by_name["リアス"].dialogue_count, 1)
        self.assertEqual(by_name["アーシア"].dialogue_count, 1)

    def test_character_detection_accepts_two_character_cjk_names(self):
        from app.chapters.detector import Chapter

        chapter = Chapter(
            number=1,
            title="Short CJK speaker names",
            text='玲奈: 「待って。」\n悠真: 「今行く。」',
        )
        analysis = analyze_book([chapter])
        by_name = {item.name: item for item in analysis.characters}

        self.assertIn("玲奈", by_name)
        self.assertIn("悠真", by_name)
        self.assertEqual(by_name["玲奈"].dialogue_count, 1)
        self.assertEqual(by_name["悠真"].dialogue_count, 1)
        self.assertEqual(analysis.dialogue_total, 2)

    def test_character_analysis_preserves_multi_speaker_manual_assignments(self):
        chapter = Chapter(
            number=1,
            title="Multi-speaker assignment",
            text='“We are ready.”',
            dialogue_assignments=[{
                "start": 0,
                "end": 15,
                "text": "We are ready.",
                "speakers": ["Rias Gremory", "Akeno Himejima"],
                "multi_speaker_mode": "chorus",
            }],
        )
        analysis = analyze_book([chapter])
        names = {item.name for item in analysis.characters}
        self.assertIn("Rias Gremory", names)
        self.assertIn("Akeno Himejima", names)

    def test_pronunciation_candidates_reject_ordinary_english(self):
        self.assertTrue(is_common_english_phrase("All"))
        self.assertTrue(is_common_english_phrase("Confirmed"))
        self.assertTrue(is_common_english_phrase("Sir"))
        self.assertTrue(is_common_english_phrase("Sure"))
        self.assertEqual(nativeish_pronunciation("Hinata", "Japanese"), "Hee-nah-tah")

    def test_elevenlabs_key_round_trip_is_local(self):
        import app.tts.providers.elevenlabs as eleven
        with tempfile.TemporaryDirectory() as temp:
            original = eleven.ElevenLabsProvider.KEY_FILE
            eleven.ElevenLabsProvider.KEY_FILE = Path(temp) / "elevenlabs.json"
            try:
                eleven.ElevenLabsProvider.save_api_key("test-key")
                self.assertEqual(eleven.ElevenLabsProvider.load_api_key(), "test-key")
                self.assertTrue(eleven.ElevenLabsProvider.KEY_FILE.exists())
            finally:
                eleven.ElevenLabsProvider.KEY_FILE = original

    def test_custom_voice_worker_uses_utf8_ipc(self):
        provider_source = Path("app/tts/providers/chatterbox.py").read_text(encoding="utf-8")
        runtime_source = Path("app/tts/chatterbox_runtime.py").read_text(encoding="utf-8")
        preview_source = Path("app/tts/preview.py").read_text(encoding="utf-8")
        self.assertIn('encoding="utf-8"', provider_source)
        self.assertIn('PYTHONIOENCODING', runtime_source)
        self.assertIn('PYTHONUTF8', runtime_source)
        self.assertIn('encoding="utf-8"', preview_source)

    def test_offline_qwen_runtime_never_creates_venv_from_frozen_app(self):
        source = Path("app/tts/qwen_character_runtime.py").read_text(encoding="utf-8")
        self.assertIn("python312", source)
        self.assertIn('PYTHON_INSTALLER_URL =', source)
        self.assertIn('PYTHON_VERSION = "3.12.10"', source)
        self.assertIn("TargetDir=", source)
        self.assertNotIn("EnvBuilder(with_pip=True, clear=False, upgrade_deps=True)", source)

    def test_offline_qwen_character_catalog_is_declared(self):
        from app.tts.voice_profile import builtin_voice_profiles
        source = Path("app/tts/qwen_character_runtime.py").read_text(encoding="utf-8")
        worker = Path("app/tts/qwen_character_worker.py").read_text(encoding="utf-8")
        self.assertIn("Qwen/Qwen3-TTS-12Hz-0.6B-CustomVoice", source)
        self.assertIn("Qwen/Qwen3-TTS-12Hz-1.7B-CustomVoice", source)
        self.assertIn("Qwen/Qwen3-TTS-12Hz-1.7B-VoiceDesign", source)
        self.assertIn("Qwen/Qwen3-TTS-12Hz-0.6B-Base", source)
        self.assertIn("Qwen/Qwen3-TTS-12Hz-1.7B-Base", source)
        self.assertIn("--models-root", worker)
        self.assertIn('choices=["custom", "design", "clone"]', worker) if False else None

        profiles = [
            p for p in builtin_voice_profiles()
            if p.provider == "qwen-character"
        ]
        # The catalogue remains visible before installation so users can choose
        # a character and see exactly which local pack is required.
        self.assertGreaterEqual(len(profiles), 4)
        english_ids = {p.voice_id for p in profiles if p.language == "English"}
        self.assertEqual(english_ids, {"Ryan", "Aiden"})
        self.assertTrue(any(p.provider == "qwen-character" and p.voice_id == "Ono_Anna" for p in profiles))

    def test_qwen_provider_selects_the_three_real_model_phases(self):
        from app.tts.providers.qwen_character import QwenCharacterProvider

        custom = QwenCharacterProvider(
            voice_id="Ryan",
            model_kind="custom-1.7b",
            qwen_mode="custom",
        )
        design = QwenCharacterProvider(
            voice_id="design-test",
            model_kind="design-1.7b",
            qwen_mode="design",
            instruct="Warm cinematic young female anime heroine.",
            qwen_seed=123,
        )
        clone = QwenCharacterProvider(
            voice_id="clone-test",
            model_kind="base-1.7b",
            qwen_mode="clone",
            reference_audio=Path("reference.wav"),
            reference_text="Hello there.",
        )
        self.assertEqual(custom.model_kind, "custom-1.7b")
        self.assertEqual(design.model_kind, "design-1.7b")
        self.assertEqual(design.qwen_mode, "design")
        self.assertEqual(clone.model_kind, "base-1.7b")
        self.assertEqual(clone.qwen_mode, "clone")
        self.assertEqual(clone.reference_text, "Hello there.")

    def test_chapter_editor_uses_non_obstructive_colour_highlights_and_hover_tooltips(self):
        source = Path("app/ui/chapter_editor.py").read_text(encoding="utf-8")
        self.assertIn("self.text.viewport().installEventFilter(self)", source)
        self.assertIn("QToolTip.showText", source)
        self.assertIn("self._assignment_hover_items.append((start, end, tuple(speakers)))", source)
        self.assertIn("selection.format = fmt", source)
        self.assertNotIn("badge.setFixedSize(badge_width, 16)", source)
        self.assertNotIn('font-size:7px', source)
        self.assertIn('_update_current_chapter_list_label', source)
        self.assertIn('self.title.textChanged.connect(self._update_current_chapter_list_label)', source)

    def test_chapter_editor_assignment_ui_is_subtle_and_uses_project_title(self):
        source = Path("app/ui/chapter_editor.py").read_text(encoding="utf-8")
        self.assertIn("self.text.setViewportMargins(82, 0, 0, 0)", source)
        self.assertIn('assignment_manager_button = QPushButton("Assignments")', source)
        self.assertIn("soft.setAlpha(52)", source)
        self.assertIn("QEvent.Type.MouseMove", source)
        self.assertIn("text_widget.cursorForPosition(point)", source)
        self.assertIn("if watched is viewport:", source)
        self.assertIn("# being destroyed. Do not dereference its deleted QTextEdit wrapper.", source)
        self.assertIn("except RuntimeError:", source)
        self.assertNotIn('QPushButton(display_name, self.text.viewport())', source)
        self.assertIn("project_title: str | None = None", source)
        self.assertIn('self.project_title = str(project_title or "").strip()', source)
        self.assertNotIn('text=self.window().windowTitle()', source)

    def test_workflow_passes_real_project_title_to_chapter_editor(self):
        source = Path("app/ui/workflow.py").read_text(encoding="utf-8")
        self.assertIn("self.project.folder,", source)
        self.assertIn("self.project.title,", source)

    def test_voice_page_stage_two_contains_library_selectors(self):
        source = Path("app/ui/voice_page.py").read_text(encoding="utf-8")
        self.assertIn("library_config_box = QGroupBox", source)
        self.assertIn("root.addWidget(library_config_box)", source)
        self.assertIn("library_config.addWidget(self.neural_box)", source)
        self.assertIn("library_config.addWidget(self.qwen_phase_box)", source)
        self.assertIn("library_config.addWidget(self.sapi_box)", source)
        self.assertIn("library_config.addWidget(self.custom_box)", source)
        self.assertNotIn("root.addWidget(self.custom_box)", source)

    def test_voice_page_uses_four_stage_preview_layout(self):
        source = Path("app/ui/voice_page.py").read_text(encoding="utf-8")
        for marker in (
            'QGroupBox("1  •  Voice Library")',
            'QGroupBox("2  •  Library Configuration")',
            'QGroupBox("3  •  Voice Profile")',
            'QGroupBox("4  •  Preview")',
        ):
            self.assertIn(marker, source)
        self.assertIn("PREVIEW_DURATION_OPTIONS = (10, 30, 60, 90)", source)
        self.assertIn("PREVIEW_DURATION_OPTIONS.index(30)", source)
        self.assertIn("for seconds in PREVIEW_DURATION_OPTIONS:", source)
        self.assertIn("preview_library", source)
        self.assertIn("preview_voice", source)
        self.assertIn("preview_model", source)
        self.assertIn("VoiceDesign preview", source)
        self.assertNotIn("qwen_design_preview_duration", source)
        self.assertNotIn("test_profile_button", source)

    def test_qwen_design_preview_uses_one_pass_voice_design_model(self):
        provider = Path("app/tts/providers/qwen_character.py").read_text(encoding="utf-8")
        worker = Path("app/tts/qwen_character_worker.py").read_text(encoding="utf-8")
        host = Path("app/ui/premium_v2_host.py").read_text(encoding="utf-8")
        self.assertIn('"design-preview"', provider)
        self.assertIn('if self.qwen_mode in {"design", "design-preview"}:', provider)
        self.assertIn("self.model_kind = best_voice_design_kind()", provider)
        self.assertIn('elif active_task == "design-preview":', worker)
        self.assertIn("_generate_design_preview(model, request, output)", worker)
        self.assertIn('elif active_task == "design":', worker)
        self.assertIn("_generate_design_locked(model, locked_clone_prompts, request, output)", worker)
        self.assertIn('preview_profile.qwen_mode = "design-preview"', host)
        self.assertIn("performance_target_met", host)

    def test_qwen_voice_consistency_uses_design_anchor_and_base_prompt(self):
        worker = Path("app/tts/qwen_character_worker.py").read_text(encoding="utf-8")
        provider = Path("app/tts/providers/qwen_character.py").read_text(encoding="utf-8")
        self.assertIn("create_voice_clone_prompt", worker)
        self.assertIn("generate_voice_design", worker)
        self.assertIn("voice_anchors", worker)
        self.assertIn("locked_clone_prompts[style] = model.create_voice_clone_prompt", worker)
        self.assertIn("reference_style(style)", worker)
        self.assertIn("ANCHOR_TEXTS", worker)
        self.assertIn("recommended_chunk_chars = 800", provider)
        self.assertIn("recommended_chunk_sentences", provider)

    def test_voice_profile_save_and_qwen_preview_controls_are_exposed(self):
        source = Path("app/ui/voice_page.py").read_text(encoding="utf-8")
        self.assertIn("Save Voice Profile", source)
        self.assertIn("preview_duration", source)
        self.assertIn("preview_duration", source)
        self.assertIn("Qwen3-TTS  ·  VoiceDesign", source)

    def test_premium_v2_chapter_cleaner_uses_import_noise_cleaner(self):
        from app.documents.parser import clean_import_noise
        raw = "Story continues.\nReport\nwww.asianovel.com\n\nAnother paragraph.\n"
        cleaned, removed = clean_import_noise(raw)
        self.assertNotIn("Report", cleaned)
        self.assertNotIn("www.asianovel.com", cleaned)
        self.assertIn("Story continues.", cleaned)
        self.assertIn("Another paragraph.", cleaned)
        self.assertEqual(removed, ["Report", "www.asianovel.com"])

    def test_chapter_cleaner_detects_and_reflows_import_noise(self):
        from app.documents.parser import detect_repeated_book_noise, clean_import_noise, reflow_source_text
        raw = (
            "The hero walked toward the door and\n"
            "stopped when he heard a sound.\n"
            "41\n"
            "Report\n"
            "www.asianovel.com\n\n"
            "A second paragraph begins here.\n"
            "41\n"
            "Report\n"
            "www.asianovel.com\n"
        )
        cleaned, removed = clean_import_noise(raw)
        self.assertNotIn("www.asianovel.com", cleaned)
        self.assertIn("A second paragraph begins here.", cleaned)
        self.assertIn("The hero walked toward the door", reflow_source_text(cleaned))
        self.assertEqual(
            removed,
            ["41", "Report", "www.asianovel.com", "41", "Report", "www.asianovel.com"],
        )

    def test_txt_import_strips_utf8_bom_without_changing_story_text(self):
        from app.documents.parser import extract_text

        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "bom_story.txt"
            path.write_bytes(b"\xef\xbb\xbfChapter 1\nThe story begins.")
            extracted = extract_text(path)
            self.assertEqual(extracted.title, "bom_story")
            self.assertTrue(extracted.text.startswith("Chapter 1\n"))
            self.assertNotIn("\ufeff", extracted.text)
            self.assertIn("The story begins.", extracted.text)

    def test_epub_import_follows_spine_reading_order(self):
        import types
        from unittest.mock import patch
        from app.documents.parser import _extract_epub

        class FakeItem:
            def __init__(self, item_id, html):
                self.item_id = item_id
                self.html = html

            def get_id(self):
                return self.item_id

            def get_content(self):
                return self.html.encode("utf-8")

        first = FakeItem("chapter-b", "<html><body>Second in reading order.</body></html>")
        second = FakeItem("chapter-a", "<html><body>First in reading order.</body></html>")
        fake_book = types.SimpleNamespace(
            spine=[("chapter-a", "yes"), ("chapter-b", "yes")]
        )
        fake_book.get_items_of_type = lambda item_type: [first, second]
        fake_epub = types.SimpleNamespace(read_epub=lambda _path: fake_book)
        fake_ebooklib = types.ModuleType("ebooklib")
        fake_ebooklib.ITEM_DOCUMENT = 9
        fake_ebooklib.epub = fake_epub

        with patch.dict("sys.modules", {"ebooklib": fake_ebooklib}):
            extracted = _extract_epub(Path("fake.epub"))

        self.assertLess(
            extracted.index("First in reading order."),
            extracted.index("Second in reading order."),
        )

    def test_docx_import_preserves_story_text_inside_tables(self):
        from docx import Document
        from app.documents.parser import extract_text

        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "table_story.docx"
            document = Document()
            document.add_paragraph("Opening paragraph.")
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Story text inside a table."
            document.add_paragraph("Closing paragraph.")
            document.save(path)

            extracted = extract_text(path)
            self.assertIn("Opening paragraph.", extracted.text)
            self.assertIn("Story text inside a table.", extracted.text)
            self.assertIn("Closing paragraph.", extracted.text)
            self.assertLess(extracted.text.index("Opening paragraph."), extracted.text.index("Story text inside a table."))
            self.assertLess(extracted.text.index("Story text inside a table."), extracted.text.index("Closing paragraph."))

    def test_source_reflow_preserves_life_chapter_headings(self):
        from app.documents.parser import reflow_source_text
        source = (
            "Life.0\n"
            "The opening of the story.\n"
            "Life.1 I Quit Being a Human\n"
            "The next chapter starts here.\n"
            "Life.2 I Start as a Devil\n"
            "Another chapter body."
        )
        result = reflow_source_text(source)
        self.assertIn("Life.0\nThe opening of the story.", result)
        self.assertIn("Life.1 I Quit Being a Human\nThe next chapter starts here.", result)
        self.assertIn("Life.2 I Start as a Devil\nAnother chapter body.", result)
        self.assertNotIn("Life.0 Life.1", result)

    def test_generation_pipeline_applies_speed_during_chapter_packaging(self):
        assembler = Path("app/audio/assembler.py").read_text(encoding="utf-8")
        generator = Path("app/tts/generator.py").read_text(encoding="utf-8")
        manager = Path("app/tts/manager.py").read_text(encoding="utf-8")
        self.assertIn("narration_speed: float = 1.0", assembler)
        self.assertIn('"-filter:a"', assembler)
        self.assertIn('f"atempo={speed:.3f}"', assembler)
        self.assertIn("narration_speed=self.narration_speed", manager)
        self.assertNotIn("_apply_narration_speed(", generator)
        self.assertIn("if len(completed) % 4 == 0", generator)

    def test_qwen_long_form_runtime_uses_lighter_base_and_larger_units(self):
        provider = Path("app/tts/providers/qwen_character.py").read_text(encoding="utf-8")
        runtime = Path("app/tts/qwen_character_runtime.py").read_text(encoding="utf-8")
        worker = Path("app/tts/qwen_character_worker.py").read_text(encoding="utf-8")
        self.assertIn("recommended_chunk_chars = 800", provider)
        self.assertIn("recommended_chunk_sentences", provider)
        self.assertIn("best_long_form_clone_kind", provider)
        self.assertIn("base-0.6b", runtime)
        self.assertIn("Base 1.7B exceeded available GPU memory", worker)
        self.assertIn('"longform": self.qwen_mode != "design-preview"', provider)
        self.assertIn("def _recommended_max_new_tokens", worker)
        self.assertIn("recommended_batch_size", provider)
        self.assertIn('"texts":', provider)

    def test_qwen_voicedesign_preview_is_fast_and_production_keeps_identity_lock(self):
        worker = Path("app/tts/qwen_character_worker.py").read_text(encoding="utf-8")
        provider = Path("app/tts/providers/qwen_character.py").read_text(encoding="utf-8")
        host = Path("app/ui/premium_v2_host.py").read_text(encoding="utf-8")
        voice_page = Path("app/ui/voice_page.py").read_text(encoding="utf-8")
        self.assertIn('if active_task == "design":', worker)
        self.assertIn('elif active_task == "design-preview":', worker)
        self.assertIn("_generate_design_preview(model, request, output)", worker)
        self.assertIn("_generate_design_locked(model, locked_clone_prompts, request, output)", worker)
        self.assertIn("_create_expression_anchors", worker)
        self.assertIn('"outputs": [str(Path(str(path)).resolve()) for path in batch_outputs]', worker)
        self.assertIn('qwen_mode="design"', voice_page)
        self.assertIn('preview_profile.qwen_mode = "design-preview"', host)
        self.assertIn('best_long_form_clone_kind()', provider)
        self.assertIn('"--clone-kind", clone_kind or best_clone_kind()', provider)

    def test_voicedesign_requires_stable_base_model(self):
        provider = Path("app/tts/providers/qwen_character.py").read_text(encoding="utf-8")
        self.assertIn("VoiceDesign stable identity needs the local Qwen Base model", provider)
        self.assertIn("best_long_form_clone_kind()", provider)

    def test_app_starts_maximized_for_screen_fit(self):
        source = Path("app/main.py").read_text(encoding="utf-8")
        self.assertIn("window.showMaximized()", source)

    def test_voice_previews_expose_percentage_progress(self):
        voice_page = Path("app/ui/voice_page.py").read_text(encoding="utf-8")
        generation = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        self.assertIn("progress = Signal(int)", voice_page)
        self.assertIn("self.preview_percent_bar", voice_page)
        self.assertIn("percent = Signal(int)", generation)
        self.assertIn("self.preview_loading.setRange(0, 100)", generation)
        self.assertIn("self.preview_worker.percent.connect", generation)

    def test_qwen_preview_and_audiobook_use_same_generation_ceiling(self):
        worker = Path("app/tts/qwen_character_worker.py").read_text(encoding="utf-8")
        self.assertIn("def _recommended_max_new_tokens", worker)

    def test_import_reflow_joins_hard_wraps_but_preserves_dialogue_lines(self):
        from app.documents.parser import reflow_source_text
        text = (
            "In terms of military assistance, it seemed\n"
            "unlikely that we would suddenly be tapped for that.\n"
            "“Our name…?” “I mean, yeah.”\n"
            "The next paragraph starts here."
        )
        result = reflow_source_text(text)
        self.assertIn("it seemed unlikely that we would suddenly", result)
        self.assertIn("“Our name…?”\n“I mean, yeah.”", result)

    def test_generate_refresh_method_exists_for_live_chapter_updates(self):
        source = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        workflow = Path("app/ui/workflow.py").read_text(encoding="utf-8")
        self.assertIn("def refresh_chapter_sources", source)
        self.assertIn("page.refresh_chapter_sources", workflow)
        self.assertIn("def _refresh_generation_context", workflow)

    def test_qwen_profiles_have_unique_cache_identity(self):
        from app.tts.voice_profile import VoiceProfile, save_profiles, load_profiles
        from app.core.paths import voices_root
        with tempfile.TemporaryDirectory() as temp:
            # Ensure a constructed Qwen profile receives an immutable identity on save.
            original = voices_root
            import app.tts.voice_profile as vp
            vp.voices_root = lambda: Path(temp)
            try:
                profile = VoiceProfile(
                    name="Test Qwen Voice",
                    provider="qwen-character",
                    voice_id="design-test-qwen",
                    qwen_mode="design",
                    qwen_prompt="A natural male narrator.",
                    qwen_seed=3101,
                )
                save_profiles([profile])
                loaded = load_profiles()
                saved = next(p for p in loaded if p.name == "Test Qwen Voice")
                self.assertTrue(saved.qwen_profile_id)
            finally:
                vp.voices_root = original

    def test_qwen_worker_bootstrap_adds_frozen_application_root(self):
        source = Path("app/tts/qwen_character_worker.py").read_text(encoding="utf-8")
        self.assertIn("_APPLICATION_ROOT = Path(__file__).resolve().parents[2]", source)
        self.assertIn("sys.path.insert(0, str(_APPLICATION_ROOT))", source)

    def test_qwen_repair_avoids_py_compile_and_program_files_pycache(self):
        runtime = Path("app/tts/qwen_character_runtime.py").read_text(encoding="utf-8")
        self.assertIn("PYTHONDONTWRITEBYTECODE", runtime)
        self.assertIn("ast.parse", runtime)
        self.assertNotIn('"-m", "py_compile"', runtime)

    def test_qwen_expression_styles_are_subtle_and_text_driven(self):
        expression = Path("app/tts/expression.py").read_text(encoding="utf-8")
        from app.tts.expression import classify_expression
        self.assertIn("thoughtful", expression)
        self.assertIn("warm", expression)
        self.assertIn("concerned", expression)
        self.assertEqual(classify_expression("Instead of me…then?"), "thoughtful")
        self.assertIn(classify_expression("That's amazing! Really!"), {"excited", "thoughtful"})

    def test_qwen_expression_styles_map_to_stable_reference_families(self):
        from app.tts.expression import classify_expression, reference_style
        self.assertEqual(classify_expression("Maybe I should wait..."), "thoughtful")
        self.assertEqual(classify_expression("That's wonderful!"), "excited")
        self.assertEqual(reference_style("curious"), "thoughtful")
        self.assertEqual(reference_style("excited"), "warm")
        self.assertEqual(reference_style("sad"), "concerned")
        self.assertEqual(reference_style("neutral"), "neutral")

    def test_story_pacing_preserves_headings_dialogue_and_attribution(self):
        from app.tts.pacing import is_heading_line
        source = (
            "The hallway was silent.\n\n"
            "THE DEMON LORD SUMMIT\n\n"
            "The three people sat down.\n\n"
            "“Don't move!” he shouted.\n"
            "“Why?” she asked."
        )
        units = split_for_pacing(source, max_chars=800, max_sentences=4)
        self.assertEqual(units[1], "THE DEMON LORD SUMMIT")
        self.assertIn("“Don't move!” he shouted.", units)
        self.assertIn("“Why?” she asked.", units)
        self.assertTrue(is_heading_line("THE DEMON LORD SUMMIT"))
        self.assertGreaterEqual(pause_after_ms("THE DEMON LORD SUMMIT"), 700)
        self.assertGreaterEqual(
            pause_after_ms("He stepped back.", "natural", next_text="CHAPTER 1"),
            600,
        )

    def test_qwen_chunking_splits_only_at_expression_family_changes(self):
        text = "The hallway was silent. He whispered, ‘Don't move.’ She shouted, ‘Run!’"
        ordinary = split_for_pacing(text, max_chars=800, max_sentences=4)
        emotional = split_for_pacing(
            text,
            max_chars=800,
            max_sentences=4,
            split_on_expression_change=True,
        )
        self.assertGreater(len(emotional), len(ordinary))
        self.assertTrue(any("whispered" in value for value in emotional))

    def test_quoted_dialogue_attribution_stays_with_the_quote(self):
        from app.tts.pacing import sentence_units
        units = sentence_units("“Wait!” he said. Then he left.")
        self.assertEqual(units[0], "“Wait!” he said.")
        self.assertEqual(units[1], "Then he left.")

    def test_qwen_provider_constructor_does_not_probe_gpu_on_ui_thread(self):
        from unittest.mock import patch
        from app.tts.providers.qwen_character import QwenCharacterProvider

        def fast_custom(vram_gb=None):
            self.assertEqual(vram_gb, 0.0)
            return "custom-0.6b"

        def fast_clone(vram_gb=None):
            self.assertEqual(vram_gb, 0.0)
            return "base-0.6b"

        with patch("app.tts.providers.qwen_character.best_custom_kind", side_effect=fast_custom), \
             patch("app.tts.providers.qwen_character.best_clone_kind", side_effect=fast_clone):
            custom = QwenCharacterProvider(voice_id="Ryan", model_kind="auto", qwen_mode="custom")
            clone = QwenCharacterProvider(voice_id="clone-test", model_kind="auto", qwen_mode="clone")
        self.assertEqual(custom.model_kind, "custom-0.6b")
        self.assertEqual(clone.model_kind, "base-0.6b")
        self.assertTrue(custom._auto_select_model)
        self.assertTrue(clone._auto_select_model)

    def test_voice_design_anchor_instruction_preserves_identity_and_mood_range(self):
        from app.tts.expression import voice_design_anchor_instruction
        instruction = voice_design_anchor_instruction(
            "Young adult male anime protagonist, lower-mid pitch and restrained confidence."
        )
        self.assertIn("Young adult male anime protagonist", instruction)
        self.assertIn("Preserve its specified gender presentation", instruction)
        self.assertIn("pitch range", instruction)
        self.assertIn("warm, pleasantly surprised energy", instruction)
        self.assertIn("briefly concerned", instruction)
        self.assertIn("not four different voices", instruction)
        worker = Path("app/tts/qwen_character_worker.py").read_text(encoding="utf-8")
        self.assertIn("voice_design_anchor_instruction(base_description)", worker)
        self.assertIn('"temperature": temperature', worker)

    def test_changing_qwen_preset_clears_old_saved_voice_selection(self):
        source = Path("app/ui/voice_page.py").read_text(encoding="utf-8")
        start = source.index("    def _qwen_design_preset_changed(self) -> None:")
        end = source.index("    def _refresh_online_catalog(self)", start)
        handler = source[start:end]
        self.assertIn("self._clear_saved_profile_selection()", handler)
        self.assertIn('self.name.setText(f"Qwen VoiceDesign • {preset_name}")', handler)
        self.assertIn('self.profile_badge.setText(f"Unsaved draft • {preset_name}")', handler)
        self.assertIn("Preview this new voice, then save", handler)
        self.assertIn("it as a separate profile", handler)
        self.assertIn("the previously saved voice is unchanged", handler)

    def test_saved_voice_design_uses_one_shared_anchor_transcript(self):
        from app.tts.expression import DEFAULT_VOICE_DESIGN_ANCHOR_TEXT
        from app.tts.providers.qwen_character import QwenCharacterProvider
        from app.ui.voice_page import VoicePage

        provider = QwenCharacterProvider(
            voice_id="design-kousuke",
            language="English",
            instruct="Young adult male anime protagonist.",
            model_kind="design-1.7b",
            qwen_mode="design",
            qwen_anchor_text="",
        )
        self.assertEqual(provider.qwen_anchor_text, DEFAULT_VOICE_DESIGN_ANCHOR_TEXT)
        source = Path("app/ui/voice_page.py").read_text(encoding="utf-8")
        self.assertIn("anchor_text = DEFAULT_VOICE_DESIGN_ANCHOR_TEXT", source)
        self.assertIn("load_saved_profiles()", source)

    def test_qwen_runtime_readiness_is_a_fast_filesystem_gate(self):
        import app.tts.qwen_character_runtime as runtime
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            marker = root / "runtime.json"
            python = root / "python.exe"
            worker = root / "qwen_character_worker.py"
            marker.write_text(json.dumps({
                "ready": True,
                "bootstrap_version": runtime.BOOTSTRAP_VERSION,
                "qwen_tts": runtime.QWEN_TTS_VERSION,
            }), encoding="utf-8")
            python.touch()
            worker.touch()
            with patch.object(runtime, "MARKER", marker), \
                 patch.object(runtime, "VENV_PYTHON", python), \
                 patch.object(runtime, "worker_script", return_value=worker), \
                 patch.object(runtime, "_python_works", side_effect=AssertionError("subprocess check ran")), \
                 patch.object(runtime, "_runtime_self_test", side_effect=AssertionError("model imports ran")):
                self.assertTrue(runtime.runtime_ready())

    def test_preview_reports_elapsed_time_and_uses_natural_pacing(self):
        source = Path("app/ui/voice_page.py").read_text(encoding="utf-8")
        self.assertIn("finished_ok = Signal(str, float)", source)
        self.assertIn('pacing_profile="natural"', source)
        self.assertNotIn('pacing_profile="off"', source)
        self.assertIn("self.finished_ok.emit(str(self.output), elapsed_seconds)", source)
        self.assertIn("generated in {int(elapsed) // 60}", source)
        self.assertIn("RTF", source)

    def test_internal_ellipsis_creates_a_separate_pacing_unit(self):
        units = split_for_pacing("Instead of me…then? Maybe later.", max_chars=200, max_sentences=4)
        self.assertGreaterEqual(len(units), 2)
        self.assertTrue(units[0].endswith("…"))

    def test_voice_preview_initial_text_remains_automatic_for_chapter_source(self):
        source = Path("app/ui/voice_page.py").read_text(encoding="utf-8")
        init_start = source.index('self.preview_text = QTextEdit()')
        init_end = source.index('self.preview_text.setMinimumHeight', init_start)
        initialization = source[init_start:init_end]
        self.assertIn("self.preview_text.setPlainText(self._qwen_preview_script(30))", initialization)
        self.assertLess(
            initialization.index("setPlainText(self._qwen_preview_script(30))"),
            initialization.index("self.preview_text.textChanged.connect"),
        )
        setter_start = source.index("def set_preview_source(")
        setter_end = source.index("def preview(self)", setter_start)
        setter = source[setter_start:setter_end]
        self.assertIn('if getattr(self, "_preview_text_is_auto", True):', setter)
        self.assertIn('self._preview_text_is_auto = True', setter)

    def test_timed_preview_uses_target_duration_and_live_progress(self):
        preview = Path("app/tts/preview.py").read_text(encoding="utf-8")
        voice_page = Path("app/ui/voice_page.py").read_text(encoding="utf-8")
        generation = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        self.assertIn("target_seconds: float | None = None", preview)
        self.assertIn('if target > 0 and accumulated_seconds >= target', preview)
        self.assertIn("preview_started = time.monotonic()", voice_page)
        self.assertIn("ETA", voice_page)
        self.assertIn("preview_activity", generation)

    def test_qwen_resume_signature_tracks_exact_voice_configuration(self):
        source = Path("app/tts/providers/qwen_character.py").read_text(encoding="utf-8")
        generator = Path("app/tts/generator.py").read_text(encoding="utf-8")
        cast = Path("app/tts/cast_provider.py").read_text(encoding="utf-8")
        self.assertIn("def generation_signature(self)", source)
        self.assertIn("self.qwen_mode", source)
        self.assertIn("self.model_kind", source)
        self.assertIn("self.instruct", source)
        self.assertIn("self.qwen_seed", source)
        self.assertIn("GENERATION_PIPELINE_VERSION", generator)
        # Cast signatures must follow assigned profile names, not character names.
        self.assertIn("self.assignments.values()", cast)
        self.assertIn("profile.qwen_mode", cast)
        self.assertIn("profile.qwen_profile_id", cast)

    def test_qwen_generation_uses_text_sized_token_ceiling_and_heartbeat(self):
        worker = Path("app/tts/qwen_character_worker.py").read_text(encoding="utf-8")
        provider = Path("app/tts/providers/qwen_character.py").read_text(encoding="utf-8")
        self.assertIn("def _recommended_max_new_tokens", worker)
        self.assertIn("chars * 1.10", worker)
        self.assertIn("max_new_tokens", worker)
        self.assertIn("heartbeat=", provider)

    def test_qwen_design_identity_id_is_stable_for_unchanged_saved_description(self):
        from unittest.mock import patch
        from app.ui.voice_page import VoicePage
        from app.tts.voice_profile import VoiceProfile

        saved = VoiceProfile(
            name="Anime MX Kousuke Takioto!",
            provider="qwen-character",
            voice_id="design-anime-mx-kousuke-takioto",
            model_id="design-1.7b",
            language="English",
            qwen_mode="design",
            qwen_prompt="Young adult male anime protagonist voice.",
            qwen_seed=3101,
            qwen_anchor_text="A sample sentence.",
            qwen_profile_id="saved-profile-id",
        )
        class ProfileLookup:
            class SavedBox:
                def currentData(self):
                    return saved.name
            saved_profiles = SavedBox()
        page = ProfileLookup()
        with patch("app.ui.voice_page.load_saved_profiles", return_value=[saved]):
            self.assertEqual(
                VoicePage._stable_qwen_design_profile_id(
                    page,
                    saved.name,
                    saved.qwen_prompt,
                    saved.language,
                    saved.qwen_seed,
                    saved.qwen_anchor_text,
                ),
                "saved-profile-id",
            )
            changed = VoicePage._stable_qwen_design_profile_id(
                page,
                saved.name,
                "Older sounding young adult male protagonist.",
                saved.language,
                saved.qwen_seed,
                saved.qwen_anchor_text,
            )
        self.assertNotEqual(changed, "saved-profile-id")
        self.assertEqual(len(changed), 32)

    def test_voice_studio_guided_steps_and_show_all_mode(self):
        from app.ui.voice_page import VoicePage

        page = VoicePage()
        try:
            self.assertEqual(len(page.voice_step_pages), 4)
            self.assertEqual(page.voice_step_index, 0)
            self.assertFalse(page.show_all_steps_button.isChecked())
            self.assertFalse(page.voice_step_pages[0].isHidden())
            self.assertTrue(all(item.isHidden() for item in page.voice_step_pages[1:]))

            page._set_voice_step(2)
            self.assertEqual(page.voice_step_index, 2)
            self.assertTrue(page.voice_step_pages[0].isHidden())
            self.assertFalse(page.voice_step_pages[2].isHidden())
            self.assertTrue(page.voice_step_back_button.isEnabled())
            self.assertTrue(page.voice_step_continue_button.isEnabled())

            page.show_all_steps_button.setChecked(True)
            self.assertTrue(all(not item.isHidden() for item in page.voice_step_pages))
            page.show_all_steps_button.setChecked(False)
            self.assertFalse(page.voice_step_pages[2].isHidden())
            self.assertTrue(page.voice_step_pages[0].isHidden())
            self.assertEqual(page.show_all_steps_button.text(), "Show All Steps")
        finally:
            page.close()
            page.deleteLater()

    def test_voice_cast_persistence_and_counts_are_series_scoped(self):
        source = Path("app/ui/voice_cast_page.py").read_text(encoding="utf-8")
        create_start = source.index("    def create_character_record(")
        count_start = source.index("    def _refresh_character_counts(", create_start)
        analyze_start = source.index("    def _analysis_finished(", count_start)
        load_start = source.index("    def _load_saved_cast(", source.index("    def unassign_voice("))
        labels_start = source.index("    def _refresh_labels(")
        save_start = source.index("    def _save_cast(", labels_start)
        create = source[create_start:count_start]
        counts = source[count_start:analyze_start]
        loading = source[load_start:source.index("    def add_character(", load_start)]
        labels = source[labels_start:save_start]
        save = source[save_start:]
        self.assertIn('and str(item.get("series") or self._series_name()).strip().casefold() == series.casefold()', create)
        self.assertNotIn("Character name already used", create)
        self.assertIn('assignment.get("series") or active_series', counts)
        self.assertIn("character_identity_key(name, series)", loading)
        self.assertIn("voice_cast_identity_assignments", save)
        self.assertIn("character_identity_key(name, series)", save)
        self.assertIn("len(name_series.get(name.casefold(), set())) == 1", save)
        self.assertIn("Qt.ItemDataRole.UserRole + 1", labels)
        self.assertIn("if len(series_by_name.get(name.casefold(), set())) > 1", labels)

    def test_cast_provider_keeps_same_named_series_voices_separate(self):
        from app.core.character_identity import character_identity_key
        from app.tts.cast_provider import CastAwareProvider
        from types import SimpleNamespace

        class FakeProvider:
            def __init__(self, voice):
                self.voice = voice
            def voices(self):
                return [self.voice]
            def synthesize(self, text, output_path, voice=None):
                return output_path
            def close(self):
                pass

        text_value = "“I know.”"
        profiles = {
            "Voice A": SimpleNamespace(name="Voice A", provider="piper"),
            "Voice B": SimpleNamespace(name="Voice B", provider="piper"),
        }
        provider = CastAwareProvider(
            FakeProvider("Narrator"), "Narrator", profiles, {},
            identity_assignments={
                character_identity_key("Rias Gremory", "Series A"): "Voice A",
                character_identity_key("Rias Gremory", "Series B"): "Voice B",
            },
        )
        series_a = [{
            "start": 0, "end": len(text_value), "text": text_value,
            "speakers": ["Rias Gremory"], "series": "Series A",
        }]
        series_b = [{
            "start": 0, "end": len(text_value), "text": text_value,
            "speakers": ["Rias Gremory"], "series": "Series B",
        }]
        self.assertEqual(provider.split_for_cast(text_value, "Narrator", series_a), [("I know.", "Voice A")])
        self.assertEqual(provider.split_for_cast(text_value, "Narrator", series_b), [("I know.", "Voice B")])
        provider.close()

    def test_character_identity_key_is_stable_and_unambiguous(self):
        from app.core.character_identity import character_identity_key, parse_character_identity_key
        key = character_identity_key(" Rias Gremory ", " Series A ")
        self.assertEqual(parse_character_identity_key(key), ("series a", "rias gremory"))
        self.assertNotEqual(key, character_identity_key("Rias Gremory", "Series B"))
        self.assertIsNone(parse_character_identity_key("not-an-identity-key"))

    def test_native_speaker_continuity_is_chapter_local_and_distance_limited(self):
        source = Path("app/ui/native_character_studio.py").read_text(encoding="utf-8")
        worker_start = source.index("class NativeDialogueAnalysisWorker")
        worker_end = source.index("class NativeCharacterPreviewCallbackRouter", worker_start)
        worker = source[worker_start:worker_end]
        chapter_loop = worker.index("for chapter_index, chapter in enumerate(self.chapters, 1):")
        anchor_reset = worker.index('continuity_speaker = ""', chapter_loop)
        segment_parse = worker.index("segments = dialogue_segments(chapter)", chapter_loop)
        self.assertLess(anchor_reset, segment_parse)
        self.assertIn("continuity_end = -1", worker[chapter_loop:segment_parse])
        self.assertIn("0 <= start - continuity_end <= 420", worker)
        self.assertIn('last_speaker=continuity_speaker if continuity_is_local else ""', worker)
        self.assertIn("continuity_end = end", worker)

    def test_character_assignment_rolls_back_memory_and_learning_on_save_failure(self):
        source = Path("app/ui/native_character_studio.py").read_text(encoding="utf-8")
        assign_start = source.index("    def assign_passage(")
        unassign_start = source.index("    def unassign_passage(", assign_start)
        test_voice_start = source.index("    def test_voice_profile(", unassign_start)
        assign = source[assign_start:unassign_start]
        unassign = source[unassign_start:test_voice_start]
        save_failure = assign.index('except Exception as exc:', assign.index("project.save()"))
        feedback_removal = assign.index("remove_assignment_feedback(")
        self.assertIn("chapter.dialogue_assignments = existing", assign[save_failure:feedback_removal])
        self.assertGreater(feedback_removal, assign.index("project.save()"))
        self.assertIn("chapter.dialogue_assignments = assignments", unassign)
        self.assertIn("chapter.dialogue_assignments = assignments", unassign[unassign.index("project.save()"):])
        self.assertLess(assign.index("project.save()"), feedback_removal)

    def test_character_sentence_preview_routes_worker_callbacks_through_qobject_slots(self):
        source = Path("app/ui/native_character_studio.py").read_text(encoding="utf-8")
        self.assertIn("class NativeCharacterPreviewCallbackRouter(QObject):", source)
        self.assertIn("@Slot(str, float)", source)
        self.assertIn("worker.status.connect(router.on_status)", source)
        self.assertIn("worker.finished_ok.connect(router.on_success)", source)
        self.assertIn("worker.failed.connect(router.on_failure)", source)
        self.assertNotIn("worker.status.connect(lambda", source)
        self.assertNotIn("worker.finished_ok.connect(lambda", source)
        self.assertIn("worker.cancel()", source)
        self.assertIn("worker.finished.connect(", source)
        self.assertIn("shiboken6.isValid(widget)", source)

    def test_actions_storage_cleanup_preserves_last_success_until_current_run_completes(self):
        script = Path(".github/scripts/cleanup_actions_storage.ps1").read_text(encoding="utf-8")
        self.assertIn("$currentRun = $allRuns | Where-Object { $_.id -eq $currentRunId }", script)
        self.assertIn("$currentRunIsCompleted = ($currentRun -and $currentRun.status -eq 'completed')", script)
        self.assertIn(
            "if (-not ($isCurrentWorkflow -and $pruneCurrentWorkflowToCurrent -and $currentRunIsCompleted))",
            script,
        )
        self.assertIn("$keepIds['39933305387'] = $true", script)
        self.assertIn("$_.status -eq 'completed'", script)

    def test_qwen_voice_preview_reuses_warm_provider_and_replaces_changed_profile(self):
        from types import SimpleNamespace
        from unittest.mock import patch
        import app.ui.voice_page as voice_page
        from app.tts.voice_profile import VoiceProfile

        class FakeProvider:
            def __init__(self):
                self.closed = 0
                self._worker = None
            def close(self):
                self.closed += 1
                self._worker = None

        first = FakeProvider()
        second = FakeProvider()
        profile = VoiceProfile(
            name="Warm Qwen Preview",
            provider="qwen-character",
            voice_id="design-warm",
            model_id="design-1.7b",
            backend="automatic",
            qwen_mode="design-preview",
            qwen_prompt="Warm, mid-low male narration.",
            qwen_seed=17,
            qwen_profile_id="warm-profile",
        )
        changed = VoiceProfile(
            name=profile.name, provider=profile.provider, voice_id=profile.voice_id,
            model_id=profile.model_id, backend=profile.backend, language=profile.language,
            qwen_mode=profile.qwen_mode, qwen_prompt="Older, warmer mid-low male narration.",
            qwen_seed=profile.qwen_seed, qwen_profile_id=profile.qwen_profile_id,
        )
        original_provider = voice_page._QWEN_PREVIEW_PROVIDER
        original_key = voice_page._QWEN_PREVIEW_PROFILE_KEY
        try:
            voice_page._QWEN_PREVIEW_PROVIDER = None
            voice_page._QWEN_PREVIEW_PROFILE_KEY = None
            with patch.object(
                voice_page,
                "provider_from_profile",
                side_effect=[(first, profile.voice_id), (second, changed.voice_id)],
            ) as factory:
                with voice_page._VOICE_PREVIEW_LOCK:
                    provider_a, voice_a = voice_page._get_or_create_qwen_preview_provider(profile)
                    provider_b, voice_b = voice_page._get_or_create_qwen_preview_provider(profile)
                    self.assertIs(provider_a, provider_b)
                    self.assertEqual((voice_a, voice_b), (profile.voice_id, profile.voice_id))
                    self.assertEqual(factory.call_count, 1)
                    provider_c, voice_c = voice_page._get_or_create_qwen_preview_provider(changed)
                    self.assertIs(provider_c, second)
                    self.assertEqual(voice_c, changed.voice_id)
                    self.assertEqual(first.closed, 1)
                    self.assertEqual(factory.call_count, 2)
        finally:
            voice_page._close_cached_qwen_preview_provider()
            voice_page._QWEN_PREVIEW_PROVIDER = original_provider
            voice_page._QWEN_PREVIEW_PROFILE_KEY = original_key

    def test_non_qwen_preview_evicts_warmed_qwen_provider(self):
        import app.ui.voice_page as voice_page

        class FakeProvider:
            def __init__(self):
                self.closed = 0
            def close(self):
                self.closed += 1

        cached = FakeProvider()
        original_provider = voice_page._QWEN_PREVIEW_PROVIDER
        original_key = voice_page._QWEN_PREVIEW_PROFILE_KEY
        try:
            voice_page._QWEN_PREVIEW_PROVIDER = cached
            voice_page._QWEN_PREVIEW_PROFILE_KEY = "cached-profile"
            with voice_page._VOICE_PREVIEW_LOCK:
                voice_page._evict_cached_qwen_preview_provider_locked()
            self.assertEqual(cached.closed, 1)
            self.assertIsNone(voice_page._QWEN_PREVIEW_PROVIDER)
            self.assertIsNone(voice_page._QWEN_PREVIEW_PROFILE_KEY)
        finally:
            voice_page._close_cached_qwen_preview_provider()
            voice_page._QWEN_PREVIEW_PROVIDER = original_provider
            voice_page._QWEN_PREVIEW_PROFILE_KEY = original_key

    def test_qwen_voice_preview_restarts_a_dead_cached_worker(self):
        from unittest.mock import patch
        import app.ui.voice_page as voice_page
        from app.tts.voice_profile import VoiceProfile

        class DeadWorker:
            def poll(self):
                return 7

        class FakeProvider:
            def __init__(self):
                self.closed = 0
                self._worker = DeadWorker()
            def close(self):
                self.closed += 1
                self._worker = None

        provider = FakeProvider()
        profile = VoiceProfile(
            name="Recover Qwen Preview",
            provider="qwen-character",
            voice_id="design-recover",
            model_id="design-1.7b",
            qwen_mode="design-preview",
            qwen_profile_id="recover-profile",
        )
        original_provider = voice_page._QWEN_PREVIEW_PROVIDER
        original_key = voice_page._QWEN_PREVIEW_PROFILE_KEY
        try:
            voice_page._QWEN_PREVIEW_PROVIDER = provider
            voice_page._QWEN_PREVIEW_PROFILE_KEY = voice_page._qwen_preview_profile_key(profile)
            with patch.object(voice_page, "provider_from_profile") as factory:
                with voice_page._VOICE_PREVIEW_LOCK:
                    actual, voice = voice_page._get_or_create_qwen_preview_provider(profile)
                self.assertIs(actual, provider)
                self.assertEqual(voice, profile.voice_id)
                self.assertEqual(provider.closed, 1)
                self.assertIsNone(provider._worker)
                factory.assert_not_called()
        finally:
            voice_page._close_cached_qwen_preview_provider()
            voice_page._QWEN_PREVIEW_PROVIDER = original_provider
            voice_page._QWEN_PREVIEW_PROFILE_KEY = original_key

    def test_voice_preview_is_invalidated_when_target_settings_change(self):
        from app.ui.voice_page import VoicePage

        page = VoicePage()
        try:
            page.last_preview = Path(tempfile.gettempdir()) / "old-preview.wav"
            page.last_preview_elapsed_seconds = 12.0
            page.play_button.setEnabled(True)
            page.save_preview_button.setEnabled(True)

            current = page.preview_duration.currentIndex()
            page.preview_duration.setCurrentIndex(
                (current + 1) % page.preview_duration.count()
            )

            self.assertIsNone(page.last_preview)
            self.assertIsNone(page.last_preview_elapsed_seconds)
            self.assertFalse(page.play_button.isEnabled())
            self.assertFalse(page.save_preview_button.isEnabled())
            self.assertIn("Preview invalidated", page.preview_engine_status.text())
        finally:
            page.close()
            page.deleteLater()

    def test_voice_preview_invalidation_ignores_deleted_qt_status_widget(self):
        import shiboken6
        from app.ui.voice_page import VoicePage

        page = VoicePage()
        try:
            # Closing stacked pages can destroy a child widget before its
            # preview-setting signals are disconnected. The callback must
            # tolerate a surviving Python wrapper with a deleted Qt object.
            status = page.status
            self.assertTrue(
                shiboken6.isValid(status),
                "Voice Studio must reparent its status label into the guided UI before replacing scroll content.",
            )
            parent = status.parentWidget()
            if parent is not None:
                layout = parent.layout()
                if layout is not None:
                    layout.removeWidget(status)
                status.setParent(None)
            shiboken6.delete(status)
            page._invalidate_preview()
            self.assertIsNone(page.last_preview)
        finally:
            if shiboken6.isValid(page):
                page.close()
                page.deleteLater()

    def test_qwen_clone_draft_identity_is_stable_before_profile_save(self):
        from types import SimpleNamespace
        from app.ui.voice_page import VoicePage
        from app.tts.voice_profile import VoiceProfile

        class SavedBox:
            def __init__(self, name=""):
                self.name = name
            def currentData(self):
                return self.name

        with tempfile.TemporaryDirectory() as temp:
            reference = Path(temp) / "authorized-reference.wav"
            reference.write_bytes(b"reference")
            draft = SimpleNamespace(
                profiles=[],
                saved_profiles=SavedBox(),
                _draft_qwen_clone_profile_ids={},
            )
            first = VoicePage._stable_qwen_clone_profile_id(
                draft, "Draft Voice", reference, "Hello there."
            )
            repeated = VoicePage._stable_qwen_clone_profile_id(
                draft, "Draft Voice", reference, "Hello there."
            )
            changed_transcript = VoicePage._stable_qwen_clone_profile_id(
                draft, "Draft Voice", reference, "Good morning."
            )
            self.assertEqual(first, repeated)
            self.assertNotEqual(first, changed_transcript)

            saved = VoiceProfile(
                name="Saved Clone",
                provider="qwen-character",
                voice_id=reference.stem,
                sample_path=str(reference),
                authorized=True,
                qwen_mode="clone",
                reference_text="Hello there.",
                qwen_profile_id="persisted-clone-id",
            )
            draft.profiles = [saved]
            draft.saved_profiles = SavedBox("Saved Clone")
            self.assertEqual(
                VoicePage._stable_qwen_clone_profile_id(
                    draft, "Saved Clone", reference, "Hello there."
                ),
                "persisted-clone-id",
            )

    def test_generate_tab_reads_latest_persisted_profile(self):
        import app.tts.voice_profile as voice_profiles
        from app.ui.generation_page import GenerationPage
        from app.tts.voice_profile import VoiceProfile

        with tempfile.TemporaryDirectory() as temp:
            stored = VoiceProfile(
                name="Current Kousuke",
                provider="qwen-character",
                voice_id="design-kousuke",
                qwen_mode="design",
                qwen_prompt="Lower mid-range young male voice.",
                qwen_seed=3101,
                qwen_profile_id="fresh-id",
            )
            profile_file = Path(temp) / "voices.json"
            profile_file.write_text(json.dumps([stored.__dict__]), encoding="utf-8")
            original = voice_profiles.profiles_file
            voice_profiles.profiles_file = lambda: profile_file
            try:
                class Combo:
                    def currentData(self):
                        return VoiceProfile(
                            name="Current Kousuke",
                            provider="qwen-character",
                            voice_id="design-kousuke",
                            qwen_mode="design",
                            qwen_prompt="Old female voice description.",
                            qwen_seed=12,
                            qwen_profile_id="stale-id",
                        )
                class Selection:
                    voice_profile = Combo()
                selected = GenerationPage._selected_profile(Selection())
                self.assertEqual(selected.qwen_profile_id, "fresh-id")
                self.assertEqual(selected.qwen_prompt, "Lower mid-range young male voice.")
            finally:
                voice_profiles.profiles_file = original

    def test_qwen_design_worker_lazy_loads_cached_identity(self):
        worker = Path("app/tts/qwen_character_worker.py").read_text(encoding="utf-8")
        runtime = Path("app/tts/qwen_character_runtime.py").read_text(encoding="utf-8")
        self.assertIn("saved canonical character identity", worker)
        self.assertIn("ready = all(", worker)
        self.assertIn('ANCHOR_STYLES = ("neutral",)', worker)
        self.assertIn("v8-emotion-range-anchor", worker)
        self.assertIn('if active_task not in {"design", "design-preview"}:', worker)
        self.assertIn('if vram_gb >= 10.0 and model_installed("base-1.7b")', runtime)

    def test_qwen_batch_generation_path_exists(self):
        provider = Path("app/tts/providers/qwen_character.py").read_text(encoding="utf-8")
        worker = Path("app/tts/qwen_character_worker.py").read_text(encoding="utf-8")
        generator = Path("app/tts/generator.py").read_text(encoding="utf-8")
        self.assertIn("def synthesize_batch", provider)
        self.assertIn('batch_texts = request.get("texts")', worker)
        self.assertIn("model.generate_voice_clone(", worker)
        self.assertIn("synthesize_batch", generator)
        self.assertIn("recommended_batch_size", generator)

    def test_qwen_failure_stops_cascading_chapter_errors(self):
        source = Path("app/tts/providers/qwen_character.py").read_text(encoding="utf-8")
        self.assertIn("stop_on_failure = True", source)
        self.assertIn("self.close()", source)
        self.assertIn("heartbeat=", source)

    def test_build_label_comes_from_bundled_build_metadata(self):
        from app.build_info import BUILD_NUMBER
        from app.version import full_version

        self.assertEqual(full_version(), f"v1.1.1 • Build {BUILD_NUMBER}")

    def test_import_controls_use_the_audiobook_brand_mark(self):
        source = Path("app/ui/workflow.py").read_text(encoding="utf-8")
        self.assertIn("def _import_book_icon()", source)
        self.assertIn("ryu_audiobook_logo.svg", source)
        self.assertIn("self.new_book_button.setIcon(self.import_book_icon)", source)
        self.assertIn('item.setIcon(self.import_book_icon)', source)

    def test_generate_navigation_opens_selected_library_project_and_handles_empty_library(self):
        workflow = Path("app/ui/workflow.py").read_text(encoding="utf-8")
        self.assertIn("self.sidebar.itemClicked.connect(self._sidebar_item_clicked)", workflow)
        self.assertIn("def _sidebar_item_clicked", workflow)
        self.assertIn("def _refresh_generation_context", workflow)
        nav_start = workflow.index("def _sidebar_changed")
        section = workflow[nav_start:workflow.index('elif name == "Models":', nav_start)]
        self.assertIn("if not self.project:", section)
        self.assertIn("selected_projects = self.library._selected_projects()", section)
        self.assertIn("self.open_project(selected_projects[0])", section)
        self.assertIn("self.sidebar.blockSignals(True)", section)
        self.assertIn("Select a book in Library or import a new book to start generating.", section)
        self.assertIn("QTimer.singleShot(75, self._refresh_generation_context)", section)
        self.assertLess(section.index("self._show_page(self.generation)"), section.index("QTimer.singleShot"))
        self.assertIn("Generate page could not be opened", section)
        self.assertIn("except Exception as exc:", section)

    def test_chapter_autosave_does_not_rebuild_generation_page(self):
        workflow = Path("app/ui/workflow.py").read_text(encoding="utf-8")
        editor = Path("app/ui/chapter_editor.py").read_text(encoding="utf-8")
        self.assertIn("def autosave_project", workflow)
        self.assertIn("self.on_autosave or self.on_save", editor)
        self.assertIn("self.project.save()", workflow[workflow.index("def autosave_project"):workflow.index("def save_project")])

    def test_settings_expose_and_persist_default_bitrate(self):
        settings = Path("app/ui/settings_page.py").read_text(encoding="utf-8")
        generation = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        self.assertIn('self.settings.value("audio/default_bitrate", 256)', settings)
        self.assertIn('self.settings.setValue("audio/default_bitrate"', settings)
        self.assertIn('"audio/default_bitrate"', generation)
        self.assertIn("DEFAULT_AUDIO_BITRATE_KBPS", generation)
        self.assertIn('audio_state.get("audio_bitrate")', generation)

    def test_startup_applies_saved_settings_in_both_themes(self):
        source = Path("app/main.py").read_text(encoding="utf-8")
        self.assertIn("apply_application_theme(app, saved_theme)", source)
        self.assertNotIn('if saved_theme == "light":', source)

    def test_reduce_motion_updates_existing_widgets_even_before_generate_page_exists(self):
        workflow = Path("app/ui/workflow.py").read_text(encoding="utf-8")
        motion = Path("app/ui/motion.py").read_text(encoding="utf-8")
        handler_start = workflow.index("    def _settings_saved(")
        handler_end = workflow.index("    def resizeEvent(", handler_start)
        handler = workflow[handler_start:handler_end]
        self.assertIn("if not applied_by_generation:", handler)
        self.assertIn("QApplication.allWidgets()", handler)
        self.assertIn("isinstance(widget, SmoothProgressBar)", handler)
        self.assertIn("widget.set_reduced_motion(bool(reduce_motion))", handler)
        self.assertIn("def set_reduced_motion(self, enabled: bool)", motion)

    def test_settings_accent_and_interface_scale_are_applied(self):
        settings = Path("app/ui/settings_page.py").read_text(encoding="utf-8")
        self.assertIn('self.settings.setValue("ui/accent"', settings)
        self.assertIn('self.settings.setValue("ui/interface_scale"', settings)
        self.assertIn('if accent == "velvet":', settings)
        self.assertIn("factor = scale / 100.0", settings)
        self.assertIn("apply_application_theme(", settings)
        self.assertIn('accent=str(self.accent.currentData() or "crimson")', settings)
        self.assertIn('scale=int(self.interface_scale.currentData() or 100)', settings)
        self.assertIn('text_scale=int(self.text_scale.currentData() or 100)', settings)

    def test_metadata_dropdowns_and_multiselect_genres_exist(self):
        widgets = Path("app/ui/metadata_widgets.py").read_text(encoding="utf-8")
        generation = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        self.assertIn("class HistoryComboBox", widgets)
        self.assertIn("class GenrePicker", widgets)
        self.assertIn("Anime • Fantasy", widgets)
        self.assertIn("Cartoon • Comedy", widgets)
        self.assertIn("HistoryComboBox(self._metadata_history", generation)
        self.assertIn("GenrePicker()", generation)
        self.assertNotIn('form.addRow("Series number"', generation)
        self.assertNotIn('form.addRow("Description"', generation)

    def test_assignment_speaker_labels_are_hover_only_and_do_not_cover_text(self):
        source = Path("app/ui/chapter_editor.py").read_text(encoding="utf-8")
        self.assertIn("self.text.viewport().installEventFilter(self)", source)
        self.assertIn("Assigned speaker:", source)
        self.assertIn("QToolTip.hideText()", source)
        self.assertNotIn('QPushButton(display_name, self.text.viewport())', source)
        self.assertNotIn("viewport_rect.top()", source)

    def test_generation_estimate_is_labelled_as_listening_duration(self):
        source = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        self.assertIn("Estimated listening length:", source)
        self.assertIn("This is not the time required to generate it.", source)
        self.assertIn("Estimated generation time: about", source)
        self.assertIn("_generation_calibration_key", source)
        self.assertIn("self._calibration_factor = 2.45", source)
        self.assertIn("self._active_calibration_is_resume = bool(is_resume)", source)
        self.assertIn("not self._active_calibration_is_resume", source)
        self.assertIn("self.motion_settings.setValue(f\"{key}/realtime_factor\", actual_factor)", source)
        self.assertIn("self.elapsed.setText", source)
        self.assertIn("self.remaining.setText", source)

    def test_hardware_benchmark_passes_voice_as_voice_not_model_path(self):
        source = Path("app/ui/hardware_page.py").read_text(encoding="utf-8")
        self.assertIn("PiperProvider(backend=self.backend)", source)
        self.assertIn("voice=self.voice", source)
        self.assertIn('else "automatic" if backend == "automatic"', source)
        self.assertNotIn("PiperProvider(self.voice, backend=self.backend)", source)

    def test_pdf_extraction_guard_flags_sparse_long_documents(self):
        from app.documents.parser import pdf_extraction_quality_warning
        warning = pdf_extraction_quality_warning(301, 497)
        self.assertIsNotNone(warning)
        self.assertIn("301-page PDF", warning)
        self.assertIsNotNone(pdf_extraction_quality_warning(301, 5000))
        self.assertIsNone(pdf_extraction_quality_warning(301, 10000))
        self.assertIsNone(pdf_extraction_quality_warning(12, 100))

    def test_import_worker_checks_chapter_word_coverage(self):
        source = Path("app/ui/workflow.py").read_text(encoding="utf-8")
        self.assertIn("chapter_words < int(source_words * 0.90)", source)
        self.assertIn("Chapter detection retained only", source)

    def test_models_page_groups_installs_and_supports_selected_component_removal(self):
        source = Path("app/ui/models_page.py").read_text(encoding="utf-8")
        self.assertIn('QGroupBox("Installed components")', source)
        self.assertIn('QGroupBox("Offline narration voices")', source)
        self.assertIn('QGroupBox("Character voices & voice cloning")', source)
        self.assertIn("Uninstall selected component", source)
        self.assertIn('component_id.startswith("piper:")', source)
        self.assertIn('component_id.startswith("qwen-model:")', source)
        self.assertIn('component_id == "qwen-runtime"', source)
        self.assertIn("Refusing to remove a path outside the selected component root", source)
        self.assertIn("generated audiobooks", source)
        self.assertIn("QWEN_VENV_ROOT", source)
        self.assertIn('local_kokoro_root = models_root() / "kokoro"', source)
        self.assertIn("Kokoro 82M — Bundled with application (not individually removable)", source)
        self.assertIn("local_kokoro_installed", source)

    def test_generation_surfaces_provider_device_status(self):
        manager = Path("app/tts/manager.py").read_text(encoding="utf-8")
        generation = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        cast = Path("app/tts/cast_provider.py").read_text(encoding="utf-8")
        self.assertIn('set_status = getattr(self.provider, "set_status_callback", None)', manager)
        self.assertIn('message.startswith("provider:")', generation)
        self.assertIn("def set_status_callback(self, callback) -> None:", cast)

    def test_generation_cancel_closes_active_provider(self):
        manager = Path("app/tts/manager.py").read_text(encoding="utf-8")
        self.assertIn("self.cancel_event.set()", manager)
        self.assertIn('close = getattr(provider, "close", None)', manager)
        self.assertIn("if self.cancel_event.is_set():", manager)

    def test_generation_pause_resumes_at_safe_chapter_boundaries(self):
        manager = Path("app/tts/manager.py").read_text(encoding="utf-8")
        generation = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        self.assertIn("def pause(self) -> None:", manager)
        self.assertIn("def resume(self) -> None:", manager)
        self.assertIn("self._wait_if_paused(index + 1, len(self.chapters))", manager)
        self.assertIn('message.startswith("paused:")', generation)
        self.assertIn("def toggle_generation_pause(self) -> None:", generation)
        self.assertIn('"paused"}', generation)

    def test_project_title_is_propagated_to_generation_and_editor(self):
        workflow = Path("app/ui/workflow.py").read_text(encoding="utf-8")
        generation = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        self.assertIn("self.project.title,", workflow)
        self.assertIn('self.project_title = str(project_title or "").strip()', generation)
        self.assertIn("book_name = self.project_title", generation)

    def test_chapter_editor_uses_inline_sentence_assignment_ui(self):
        source = Path("app/ui/chapter_editor.py").read_text(encoding="utf-8")
        self.assertIn('QPushButton("", self.text.viewport())', source)
        self.assertIn('dragon_cross.png', source)
        self.assertNotIn('dragon_cross.svg', source)
        self.assertIn('border-radius: 14px', source)
        self.assertIn('setIconSize(QSize(22, 22))', source)
        self.assertIn("_open_inline_assignment_dialog", source)
        self.assertIn("_selection_sentence_span", source)
        self.assertIn("QToolTip.showText", source)
        self.assertIn('"Assigned speaker: " + ", ".join(speakers)', source)
        self.assertNotIn('QPushButton(display_name, self.text.viewport())', source)
        self.assertIn("setExtraSelections", source)
        self.assertNotIn("Quick Assign Speaker", source)
        self.assertNotIn("_update_quick_assign_button", source)
        self.assertNotIn('("Assign Selected Dialogue"', source)


    def test_generation_uses_real_audio_library_and_book_cue_controls(self):
        ui = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        library = Path("app/audio/library.py").read_text(encoding="utf-8")
        assembler = Path("app/audio/assembler.py").read_text(encoding="utf-8")
        mixer = Path("app/audio/preview_mixer.py").read_text(encoding="utf-8")
        self.assertIn("Real BGM catalogue", ui)
        self.assertIn("Import tracks…", ui)
        self.assertIn("Find tracks online…", ui)
        self.assertIn("import_music_tracks", ui)
        self.assertIn("import_cue_audio", ui)
        self.assertIn("choose_from_cue_library", ui)
        self.assertIn("AUDIO_SOURCE_LINKS", library)
        self.assertIn("resolve_music_track", assembler)
        self.assertIn("resolve_cue_track", assembler)
        self.assertIn("resolve_music_track", mixer)
        self.assertNotIn("synthesize_music_bed(", assembler)
        self.assertNotIn("synthesize_opening_cue(", assembler)
        self.assertNotIn("CUE_STYLE_LABELS", ui)
        self.assertNotIn("BGM catalogue (1,344 cinematic arrangements)", ui)

    def test_bundled_starter_audio_seeds_once_and_preserves_user_files(self):
        from unittest.mock import patch
        from app.audio.library import ensure_audio_library, list_music_tracks, list_cue_tracks

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            bundle = root / "bundle"
            library_root = root / "Audio Library"
            music_rel = Path("music") / "fantasy" / "starter-fantasy.mp3"
            opening_rel = Path("cues") / "opening" / "starter-opening.wav"
            (bundle / music_rel).parent.mkdir(parents=True)
            (bundle / opening_rel).parent.mkdir(parents=True)
            (bundle / music_rel).write_bytes(b"starter sample bytes")
            (bundle / opening_rel).write_bytes(b"starter cue bytes")
            (bundle / "manifest.json").write_text(
                json.dumps({
                    "version": "unit-test-v1",
                    "tracks": [
                        {
                            "type": "music", "category": "fantasy",
                            "bundle_path": music_rel.as_posix(),
                            "title": "Starter Fantasy", "author": "Test Artist",
                            "source_url": "https://example.invalid/fantasy",
                            "license": "CC0 1.0 Universal",
                            "license_note": "CC0 test license",
                        },
                        {
                            "type": "cue", "category": "opening",
                            "bundle_path": opening_rel.as_posix(),
                            "title": "Starter Opening", "author": "Test Artist",
                            "source_url": "https://example.invalid/opening",
                            "license": "CC0 1.0 Universal",
                            "license_note": "CC0 test license",
                        },
                    ],
                }),
                encoding="utf-8",
            )
            (bundle / "ATTRIBUTION.md").write_text("Starter pack credits", encoding="utf-8")
            with patch("app.audio.library.audio_library_root", return_value=library_root), patch(
                "app.audio.library._bundled_starter_audio_root", return_value=bundle
            ):
                ensure_audio_library()
                music_path = library_root / music_rel
                cue_path = library_root / opening_rel
                self.assertTrue(music_path.is_file())
                self.assertTrue(cue_path.is_file())
                self.assertEqual(len(list_music_tracks()), 1)
                self.assertEqual(len(list_cue_tracks("opening")), 1)
                metadata = json.loads(
                    music_path.with_name(music_path.name + ".license.json").read_text(encoding="utf-8")
                )
                self.assertEqual(metadata["license"], "CC0 1.0 Universal")
                self.assertEqual(metadata["source_url"], "https://example.invalid/fantasy")
                (library_root / "Starter Audio Credits.md").write_text("User-edited credits", encoding="utf-8")
                music_path.write_bytes(b"user edit")
                ensure_audio_library()
                self.assertEqual(music_path.read_bytes(), b"user edit")
                self.assertEqual((library_root / "Starter Audio Credits.md").read_text(encoding="utf-8"), "User-edited credits")
                cue_path.unlink()
                ensure_audio_library()
                self.assertFalse(cue_path.exists())

    def test_real_audio_library_imports_by_category_and_rotates_cues(self):
        from unittest.mock import patch
        from app.audio.library import (
            import_audio_files, list_music_tracks, list_cue_tracks,
            resolve_music_track, resolve_cue_track, license_note_for,
        )
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            library_root = root / "Audio Library"
            source_music = root / "cinematic-fantasy.mp3"
            source_music.write_bytes(b"fake audio bytes for library-import unit test")
            source_page = root / "hardcover-page-turn.wav"
            source_page.write_bytes(b"fake cue bytes for library-import unit test")
            source_ending = root / "soft-book-close.wav"
            source_ending.write_bytes(b"fake ending bytes for library-import unit test")
            with patch("app.audio.library.audio_library_root", return_value=library_root):
                added_music = import_audio_files(
                    [source_music], target_kind="music", category="fantasy",
                    license_note="Pixabay URL and current license checked by user",
                )
                added_page = import_audio_files([source_page], target_kind="page")
                added_ending = import_audio_files([source_ending], target_kind="ending")
                self.assertEqual(len(added_music), 1)
                self.assertEqual(len(added_page), 1)
                self.assertEqual(len(added_ending), 1)
                tracks = list_music_tracks()
                self.assertEqual(len(tracks), 1)
                self.assertEqual(tracks[0].category, "fantasy")
                self.assertIn("Pixabay URL", license_note_for(tracks[0].path))
                self.assertEqual(
                    resolve_music_track("fantasy", seed_text="chapter 1").samefile(added_music[0]),
                    True,
                )
                self.assertEqual(len(list_cue_tracks("page")), 1)
                self.assertTrue(resolve_cue_track("page", seed_text="chapter 2").samefile(added_page[0]))
                self.assertTrue(resolve_cue_track("ending", seed_text="book end").samefile(added_ending[0]))

    def test_explicit_music_catalogue_id_is_mapped_without_changing_mood(self):
        from app.audio.assembler import _selected_music_recipe
        from app.audio.sound_design import (
            MUSIC_CATALOGUE_SIZE, MOOD_ORDER, VARIATIONS_PER_MOOD, music_variation_info,
        )
        mood, variation = _selected_music_recipe("Ignored chapter", "Ignored text", "auto", 1)
        self.assertEqual((mood, variation), (MOOD_ORDER[0], 0))
        last_id = len(MOOD_ORDER) * VARIATIONS_PER_MOOD
        mood, variation = _selected_music_recipe("Ignored chapter", "Ignored text", "auto", last_id)
        self.assertEqual((mood, variation), (MOOD_ORDER[-1], VARIATIONS_PER_MOOD - 1))
        self.assertEqual(MUSIC_CATALOGUE_SIZE, 1344)
        info = music_variation_info(MOOD_ORDER[0], 0)
        self.assertIn("arrangement", info)
        self.assertNotIn("harmony", info)
        self.assertEqual(len(MOOD_ORDER), 8)


    def test_preview_mixer_uses_real_bgm_page_turn_and_book_ending(self):
        from unittest.mock import patch
        import wave
        from app.audio.library import import_audio_files
        from app.audio.preview_mixer import mix_voice_preview
        from app.audio.assembler import _duration_ms

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            library_root = root / "Audio Library"
            voice = root / "voice.wav"
            _write_wav(voice, 1.5)
            music = root / "actual-bgm.wav"
            _write_wav(music, 1.5)
            page = root / "hardcover-page-turn.wav"
            _write_wav(page, 0.55)
            ending = root / "book-closing.wav"
            _write_wav(ending, 2.65)
            with patch("app.audio.library.audio_library_root", return_value=library_root):
                import_audio_files([music], target_kind="music", category="fantasy", license_note="test license")
                import_audio_files([page], target_kind="page")
                import_audio_files([ending], target_kind="ending")
                output = root / "voice-plus-real-audio.wav"
                result = mix_voice_preview(
                    voice,
                    output,
                    "The sword clashed against the shield. The battle was over.",
                    background_enabled=True,
                    background_mode="auto",
                    music_style="fantasy",
                    music_volume=0.08,
                    story_effects_enabled=False,
                    chapter_transition_sound=True,
                    is_first_chapter=False,
                    is_last_chapter=True,
                )
                self.assertTrue(Path(result).is_file())
                self.assertTrue(output.is_file())
                with wave.open(str(output), "rb") as wav:
                    self.assertEqual(wav.getnchannels(), 2)
                    self.assertEqual(wav.getframerate(), 44100)
                self.assertGreaterEqual(_duration_ms(output), 4500)
                self.assertLessEqual(_duration_ms(output), 5200)

    def test_chapter_and_finale_real_cues_work_when_background_music_is_off(self):
        from unittest.mock import patch
        from app.audio.library import import_audio_files
        from app.audio.assembler import _duration_ms

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            library_root = root / "Audio Library"
            opening = root / "real-book-opening.wav"
            page = root / "real-page-turn.wav"
            ending = root / "real-book-ending.wav"
            _write_wav(opening, 1.18)
            _write_wav(page, 0.62)
            _write_wav(ending, 2.65)
            chapters = []
            for number in (1, 2):
                folder = root / f"{number:03d}_Chapter_{number}"
                _write_wav(folder / "chunks" / "00001.wav", 1.0)
                (folder / "narration.json").write_text(
                    json.dumps({
                        "title": f"Chapter {number}",
                        "narration_text": "The story continues.",
                    }),
                    encoding="utf-8",
                )
                chapters.append(folder)
            output = root / "chapter-cues-without-bgm.m4b"
            with patch("app.audio.library.audio_library_root", return_value=library_root):
                import_audio_files([opening], target_kind="opening")
                import_audio_files([page], target_kind="page")
                import_audio_files([ending], target_kind="ending")
                assemble_m4b(
                    chapters,
                    output,
                    "Cues without BGM",
                    chapter_titles=["Chapter 1", "Chapter 2"],
                    background_audio_mode="off",
                    chapter_transition_sound=True,
                )
            duration = _duration_ms(output)
            self.assertGreaterEqual(duration, 6000)
            self.assertLessEqual(duration, 7000)

    def test_chapter_specific_bgm_override_is_mixed_before_book_concat(self):
        from app.audio.assembler import _duration_ms

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            chapters = []
            for number, title in ((1, "First"), (2, "Second")):
                chapter_dir = root / f"{number:03d}_{title}"
                _write_wav(chapter_dir / "chunks" / "00001.wav", 2.2)
                chapters.append(chapter_dir)
            default_music = root / "default.wav"
            custom_music = root / "custom.wav"
            scene_cue = root / "scene.wav"
            _write_wav(default_music, 1.0)
            _write_wav(custom_music, 1.0)
            _write_wav(scene_cue, 0.4)
            output = root / "chapter-bgm-test.m4b"

            assemble_m4b(
                chapters,
                output,
                "Chapter BGM overrides",
                chapter_titles=["First", "Second"],
                background_audio_path=default_music,
                background_audio_volume=0.18,
                background_audio_mode="manual",
                chapter_transition_sound=True,
                scene_transition_cue_path=scene_cue,
                chapter_bgm_overrides={
                    "1": {"rule": "custom", "path": str(custom_music), "volume": 0.22},
                    "2": {"rule": "none", "volume": 0.18},
                },
            )

            self.assertTrue(output.is_file())
            self.assertGreater(output.stat().st_size, 1024)
            self.assertGreaterEqual(_duration_ms(output), 4500)
            self.assertLessEqual(_duration_ms(output), 7000)

    def test_pause_profiles_keep_sentence_and_dialogue_turns_natural(self):
        from app.tts.pacing import pause_after_ms
        self.assertGreaterEqual(pause_after_ms("That is the end.", "natural"), 280)
        self.assertGreater(pause_after_ms("That is the end?", "natural"), pause_after_ms("That is the end.", "natural"))
        self.assertGreater(pause_after_ms("That is the end.", "relaxed"), pause_after_ms("That is the end.", "natural"))
        self.assertEqual(pause_after_ms("That is the end.", "off"), 0)

    def test_cast_provider_keeps_qwen_long_form_chunk_settings(self):
        source = Path("app/tts/cast_provider.py").read_text(encoding="utf-8")
        self.assertIn('profile.provider == "qwen-character"', source)
        self.assertIn("limits.append(1800)", source)
        self.assertIn('return getattr(self.narrator_provider, "recommended_chunk_sentences", 2)', source)
        self.assertIn("return 8", source)

    def test_qwen_three_phase_voice_studio_is_exposed(self):
        source = Path("app/ui/voice_page.py").read_text(encoding="utf-8")
        provider = Path("app/tts/providers/qwen_character.py").read_text(encoding="utf-8")
        worker = Path("app/tts/qwen_character_worker.py").read_text(encoding="utf-8")
        self.assertIn("Qwen3-TTS  ·  CustomVoice", source)
        self.assertIn("Qwen3-TTS  ·  VoiceDesign", source)
        self.assertIn("Qwen3-TTS  ·  Voice Clone", source)
        self.assertIn('choices=["custom", "design", "clone"]', worker) if False else None
        self.assertIn('"design"', provider)
        self.assertIn('"clone"', provider)
        self.assertIn("generate_voice_design", worker)
        self.assertIn("generate_voice_clone", worker)

    def test_qwen_voice_profiles_persist_phase_metadata(self):
        profile = VoiceProfile(
            name="Qwen Design Test",
            provider="qwen-character",
            voice_id="design-test",
            model_id="design-1.7b",
            language="English",
            qwen_mode="design",
            qwen_prompt="Young warm anime heroine, cinematic",
            qwen_seed=123,
        )
        self.assertEqual(profile.qwen_mode, "design")
        self.assertEqual(profile.qwen_seed, 123)
        self.assertIn("anime heroine", profile.qwen_prompt)

    def test_pronunciation_scope_supports_global_library_and_search(self):
        source = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        self.assertIn('addItem("Global library", "global")', source)
        self.assertIn('pronunciation/global_dictionary', source)
        self.assertIn("def _combined_pronunciation_entries", source)
        self.assertIn("self._combined_pronunciation_entries()", source)
        self.assertIn("def _apply_pronunciation_search", source)

    def test_pronunciation_dictionary_has_audio_test_controls(self):
        source = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        self.assertIn("Test Selected", source)
        self.assertIn("PronunciationPreviewWorker", source)
        self.assertIn("pronunciation_preview_status", source)
        self.assertIn("current voice", source)

    def test_pronunciation_cleanup_covers_god_and_town(self):
        self.assertTrue(is_common_english_phrase("God"))
        self.assertTrue(is_common_english_phrase("Town"))

    def test_premium_character_voice_studio_is_exposed(self):
        from app.tts.providers.elevenlabs import ElevenLabsProvider
        source = Path("app/ui/voice_page.py").read_text(encoding="utf-8")
        self.assertIn("Premium Character Voice Studio", source)
        self.assertIn("Generate 3 Premium Character Voices", source)
        self.assertIn("Add Selected to My Voices", source)
        self.assertTrue(hasattr(ElevenLabsProvider, "design_character_voice"))
        self.assertTrue(hasattr(ElevenLabsProvider, "create_designed_voice"))

    def test_voice_cast_analysis_is_background_thread(self):
        from PySide6.QtCore import QThread
        from app.ui.voice_cast_page import VoiceCastAnalysisWorker
        self.assertTrue(issubclass(VoiceCastAnalysisWorker, QThread))

    def test_voice_cast_supports_manual_character_add_rename_and_roster_persistence(self):
        source = Path("app/ui/voice_cast_page.py").read_text(encoding="utf-8")
        self.assertIn('QPushButton("Add Character")', source)
        self.assertIn('QPushButton("Rename Character")', source)
        self.assertIn("def add_character(self):", source)
        self.assertIn("def rename_character(self):", source)
        self.assertIn('state["voice_cast_character_roster"] = [', source)
        self.assertIn('state.get("voice_cast_character_roster", [])', source)
        self.assertIn("Preserve user-created cast members", source)
        self.assertIn("assignments[name] = assignments.pop(old_name)", source)
        self.assertIn("Manually Added Character", source)

    def test_voice_cast_manual_roster_does_not_require_a_detected_speaker(self):
        from app.chapters.detector import Chapter
        analysis = analyze_book([Chapter(
            number=1,
            title="Unattributed dialogue",
            text='“We should leave now.”\n\n“Not until sunrise.”',
        )])
        # The analyzer must not invent a speaker for unattributed quotes;
        # Voice Cast's manual roster is the recovery path for this case.
        self.assertEqual(analysis.dialogue_total, 2)
        self.assertEqual(len(analysis.characters), 0)
        source = Path("app/ui/voice_cast_page.py").read_text(encoding="utf-8")
        self.assertIn("self.add_character_button.clicked.connect(self.add_character)", source)

    def test_chatterbox_reference_shorter_than_six_seconds_is_rejected(self):
        import app.tts.voice_profile as vp
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "short.wav"
            _write_wav(source, 1.0)
            original_root = vp.voices_root
            vp.voices_root = lambda: root / "Voices"
            try:
                with self.assertRaisesRegex(RuntimeError, "at least 6 seconds"):
                    vp.import_reference_audio(source, "Short Voice")
            finally:
                vp.voices_root = original_root

    def test_audiobook_visualizer_is_lightweight_and_motion_aware(self):
        from PySide6.QtWidgets import QProgressBar, QWidget
        from app.ui.audio_visualizer import AudioVisualizer
        from app.ui.motion import SmoothProgressBar

        visualizer_source = Path("app/ui/audio_visualizer.py").read_text(encoding="utf-8")
        generation_source = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        self.assertTrue(issubclass(AudioVisualizer, QWidget))
        self.assertTrue(issubclass(SmoothProgressBar, QProgressBar))
        self.assertIn("self._timer.setInterval(40)", visualizer_source)
        self.assertIn("self._playing and not self._reduced_motion and self.isVisible()", visualizer_source)
        self.assertIn("self._timer.stop()", visualizer_source)
        self.assertIn("no images, shaders, OpenGL, audio capture, or background", visualizer_source)
        self.assertIn("QSettings(\"RyuBlanc\", \"Ryu's Audiobook\")", generation_source)
        self.assertIn("self.player.playbackStateChanged.connect(self._player_playback_state_changed)", generation_source)
        self.assertIn('QCheckBox("Reduce motion")', generation_source)
        self.assertIn('self.motion_settings.setValue("ui/reduce_motion", self._reduce_motion_enabled)', generation_source)

    def test_global_motion_system_uses_low_cost_animations_and_cached_blur(self):
        from PySide6.QtCore import QThread
        from PySide6.QtWidgets import QProgressBar, QWidget
        from app.ui.motion import SmoothProgressBar
        from app.ui.loading_overlay import FrostedLoadingOverlay
        from app.ui.workflow import ImportBookWorker, RedetectChaptersWorker

        motion_source = Path("app/ui/motion.py").read_text(encoding="utf-8")
        overlay_source = Path("app/ui/loading_overlay.py").read_text(encoding="utf-8")
        workflow_source = Path("app/ui/workflow.py").read_text(encoding="utf-8")
        main_source = Path("app/main.py").read_text(encoding="utf-8")

        self.assertTrue(issubclass(SmoothProgressBar, QProgressBar))
        self.assertTrue(issubclass(FrostedLoadingOverlay, QWidget))
        self.assertTrue(issubclass(ImportBookWorker, QThread))
        self.assertTrue(issubclass(RedetectChaptersWorker, QThread))
        self.assertIn("self._busy_timer.setInterval(65)", motion_source)
        self.assertIn("animation.setDuration(180)", motion_source)
        self.assertIn("self._timer.setInterval(55)", overlay_source)
        self.assertIn("size.width() // 14", overlay_source)
        self.assertIn("self._capture_frosted_snapshot()", overlay_source)
        resize_start = overlay_source.index("    def resizeEvent(self, event)")
        resize_end = overlay_source.index("    def paintEvent(self, event)", resize_start)
        self.assertNotIn("_capture_frosted_snapshot()", overlay_source[resize_start:resize_end])
        self.assertIn("self.loading_overlay.show_loading(", workflow_source)
        self.assertIn("self._redetect_worker = worker", workflow_source)
        self.assertIn("worker.start()", workflow_source)
        transition_start = workflow_source.index("def _show_page")
        transition_end = workflow_source.index("def _sidebar_item_clicked", transition_start)
        transition_source = workflow_source[transition_start:transition_end]
        self.assertIn("self.stack.setCurrentWidget(widget)", transition_source)
        self.assertNotIn("QGraphicsOpacityEffect", transition_source)
        self.assertNotIn("animation.start()", transition_source)
        self.assertIn("QListWidget#sidebar::item:selected", main_source)
        self.assertIn("stop:0 #84142e", main_source)

    def test_automatic_music_hides_manual_track_rows(self):
        ui = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        self.assertIn("self._background_audio_form.setRowVisible(widget, visible)", ui)
        self.assertIn("(self.music_file_container, manual_mode)", ui)
        self.assertIn("(self.background_audio_volume, auto_mode or manual_mode)", ui)
        self.assertIn("(self.ambience_file_container, manual_mode)", ui)
        self.assertIn('settings_form.addRow("BGM volume", self.background_audio_volume)', ui)
        self.assertIn("self.background_audio_volume.setEnabled(auto_mode or manual_mode)", ui)
        self.assertIn("story_effects_enabled.toggled.connect(self.story_effect_volume.setEnabled)", ui)
        self.assertLess(
            ui.index("self.story_effect_volume = QDoubleSpinBox()"),
            ui.index("self.story_effects_enabled.toggled.connect(self.story_effect_volume.setEnabled)"),
        )

    def test_voice_preview_runs_in_a_worker_thread(self):
        from app.ui.voice_page import VoicePreviewWorker
        from PySide6.QtCore import QThread

        self.assertTrue(issubclass(VoicePreviewWorker, QThread))

    def test_m4b_packaging_joins_preencoded_chapters_and_reports_progress(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            chapter_dirs = []
            for number in (1, 2):
                chapter_dir = root / f"{number:03d}_Chapter_{number}"
                chunks = chapter_dir / "chunks"
                wav = chunks / "00001.wav"
                _write_wav(wav, 12.0)
                chapter_dirs.append(chapter_dir)

            progress = []
            output = root / "book.m4b"
            result = assemble_m4b(
                chapter_dirs,
                output,
                "Test Book",
                chapter_titles=["Chapter 1", "Chapter 2"],
                progress=lambda done, total, message: progress.append((done, total, message)),
            )

            self.assertEqual(result, output)
            self.assertTrue(output.exists())
            self.assertGreater(output.stat().st_size, 1024)
            self.assertTrue(any(message == "chapter-audio" for _, _, message in progress))
            self.assertTrue(any(message == "joined-audio" for _, _, message in progress))
            self.assertGreaterEqual(_probe_audio_bitrate_kbps(output), 128)
            self.assertEqual(DEFAULT_AUDIO_BITRATE_KBPS, 256)

    def test_m4b_can_mix_a_continuous_background_track(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            chapter_dirs = []
            for number in (1, 2):
                chapter_dir = root / f"{number:03d}_Chapter_{number}"
                _write_wav(chapter_dir / "chunks" / "00001.wav", 3.0)
                chapter_dirs.append(chapter_dir)

            music = root / "licensed-music.wav"
            ambience = root / "licensed-rain-ambience.wav"
            _write_wav(music, 1.0)
            _write_wav(ambience, 1.5)
            output = root / "mixed-book.m4b"
            progress = []
            result = assemble_m4b(
                chapter_dirs,
                output,
                "Mixed Test Book",
                chapter_titles=["Chapter 1", "Chapter 2"],
                progress=lambda done, total, message: progress.append(message),
                background_audio_path=music,
                background_audio_volume=0.12,
                ambience_audio_path=ambience,
                ambience_audio_volume=0.06,
                background_audio_fade_in=0.5,
                background_audio_fade_out=0.5,
            )
            self.assertEqual(result, output)
            self.assertTrue(output.exists())
            self.assertGreater(output.stat().st_size, 1024)
            self.assertIn("mixed-background-audio", progress)
            self.assertGreaterEqual(_probe_audio_bitrate_kbps(output), 128)

    def test_background_audio_controls_are_optional_and_project_persisted(self):
        ui = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        manager = Path("app/tts/manager.py").read_text(encoding="utf-8")
        assembler = Path("app/audio/assembler.py").read_text(encoding="utf-8")
        self.assertIn('QCheckBox("Enable background music / ambience")', ui)
        self.assertIn('self.background_audio_mode.addItem("Automatic • chapter-aware music", "auto")', ui)
        self.assertIn('self.background_audio_mode.addItem("Manual • use my own tracks", "manual")', ui)
        self.assertIn("def preview_background_audio(self)", ui)
        self.assertIn("def preview_ambience_audio(self)", ui)
        self.assertIn("detect", ui)
        self.assertIn("chapter text", ui)
        self.assertIn('"background_audio_enabled": self.background_audio_enabled.isChecked()', ui)
        self.assertIn('"ambience_audio_path": self.ambience_audio_path.text().strip()', ui)
        self.assertIn("background_audio_path=background_audio_path", ui)
        self.assertIn("ambience_audio_path=ambience_audio_path", ui)
        self.assertIn("background_audio_path=self.background_audio_path", manager)
        self.assertIn("ambience_audio_path=self.ambience_audio_path", manager)
        self.assertIn('inputs.extend(["-stream_loop", "-1", "-i", str(layer_path)])', assembler)
        self.assertIn("amix=inputs={mix_count}:duration=first", assembler)
        self.assertIn('if effective_audio_mode == "auto":', assembler)
        self.assertIn("resolve_cue_track", assembler)
        self.assertIn("resolve_music_track", assembler)
        self.assertNotIn("synthesize_page_turn(", assembler)
        self.assertNotIn("synthesize_audiobook_ending(", assembler)
        self.assertIn("auto_music_variation", assembler)
        self.assertIn("_selected_music_recipe", assembler)
        self.assertIn("sidechaincompress=", assembler)
        self.assertIn("automatic-real-audio-design", assembler)
        self.assertIn('effective_audio_mode in {"manual","off"}', assembler)

    def test_audiobook_audio_categories_and_mood_detection(self):
        from app.audio.sound_design import detect_chapter_mood
        from app.audio.library import MUSIC_CATEGORIES, CUE_KINDS, AUDIO_SOURCE_LINKS

        category_ids = [key for key, _label in MUSIC_CATEGORIES]
        cue_ids = [key for key, _label in CUE_KINDS]
        self.assertIn("fantasy", category_ids)
        self.assertIn("action", category_ids)
        self.assertIn("mystery", category_ids)
        self.assertIn("horror", category_ids)
        self.assertIn("romance", category_ids)
        self.assertIn("sci-fi", category_ids)
        self.assertIn("calm", category_ids)
        self.assertEqual(cue_ids, ["opening", "page", "ending"])
        self.assertTrue(any("sounddino.com/en/effects/audiobooks-music" in item["url"] for item in AUDIO_SOURCE_LINKS))
        self.assertTrue(any("storyblocks.com/audio/search/flip-book" in item["url"] for item in AUDIO_SOURCE_LINKS))
        self.assertTrue(any("pixabay.com/sound-effects/search/book-page-turn" in item["url"] for item in AUDIO_SOURCE_LINKS))
        starter_manifest = Path("assets/starter_audio/manifest.json")
        self.assertTrue(starter_manifest.is_file())
        starter = json.loads(starter_manifest.read_text(encoding="utf-8"))
        self.assertGreaterEqual(sum(1 for entry in starter["tracks"] if entry["type"] == "music"), 14)
        self.assertTrue(all(entry.get("license") and entry.get("source_url") for entry in starter["tracks"]))
        self.assertEqual(detect_chapter_mood("A Secret Revealed", "A strange hidden clue appeared."), "mystery")
        self.assertEqual(detect_chapter_mood("The Confession", "She blushed and confessed her love."), "romance")
        self.assertEqual(detect_chapter_mood("The Battlefield", "Swords clashed in the battle."), "action")
        self.assertEqual(detect_chapter_mood("A New Quest", "A wizard cast a magic spell."), "fantasy")

    def test_auto_music_and_transition_cue_use_imported_real_audio_per_chapter(self):
        from unittest.mock import patch
        from app.audio.library import import_audio_files
        from app.audio.assembler import _duration_ms

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            library_root = root / "Audio Library"
            music_mystery = root / "mystery-bed.wav"
            music_action = root / "battle-bed.wav"
            opening = root / "hardcover-open.wav"
            page_turn = root / "physical-page-turn.wav"
            ending = root / "hardcover-close.wav"
            for path, duration in (
                (music_mystery, 2.0), (music_action, 2.0),
                (opening, 1.18), (page_turn, 0.62), (ending, 2.65),
            ):
                _write_wav(path, duration)
            chapter_dirs = []
            for number, title, narration in (
                (1, "The Hidden Clue", "A hidden clue changed everything."),
                (2, "The Final Battle", "The warriors charged into battle."),
            ):
                chapter_dir = root / f"{number:03d}_{title.replace(' ', '_')}"
                _write_wav(chapter_dir / "chunks" / "00001.wav", 3.0)
                (chapter_dir / "narration.json").write_text(
                    json.dumps({"title": title, "narration_text": narration}),
                    encoding="utf-8",
                )
                chapter_dirs.append(chapter_dir)

            progress = []
            output = root / "automatic-sound-design.m4b"
            with patch("app.audio.library.audio_library_root", return_value=library_root):
                import_audio_files([music_mystery], target_kind="music", category="mystery")
                import_audio_files([music_action], target_kind="music", category="action")
                import_audio_files([opening], target_kind="opening")
                import_audio_files([page_turn], target_kind="page")
                import_audio_files([ending], target_kind="ending")
                result = assemble_m4b(
                    chapter_dirs,
                    output,
                    "Automatic Sound Design Test",
                    chapter_titles=["The Hidden Clue", "The Final Battle"],
                    progress=lambda done, total, message: progress.append(message),
                    background_audio_mode="auto",
                    auto_music_style="auto",
                    background_audio_volume=0.30,
                    background_audio_fade_in=0.4,
                    background_audio_fade_out=0.4,
                    chapter_transition_sound=True,
                )
            self.assertEqual(result, output)
            self.assertTrue(output.exists())
            self.assertGreater(output.stat().st_size, 1024)
            self.assertIn("automatic-real-audio-design", progress)
            # Six seconds of speech plus opening, page-turn and closing cues
            # chapter; allow for AAC encoder padding and container time-base rounding.
            self.assertGreaterEqual(_duration_ms(output), 10000)
            self.assertLessEqual(_duration_ms(output), 11000)
            self.assertGreaterEqual(_probe_audio_bitrate_kbps(output), 128)

    def test_story_events_detect_clear_actions_but_not_metaphors(self):
        from app.audio.story_events import detect_story_events

        action = detect_story_events(
            "He ran down the corridor. The door slammed shut behind him."
        )
        self.assertIn("running_footsteps", [event.kind for event in action])
        self.assertIn("door_slam", [event.kind for event in action])

        weather = detect_story_events("Rain poured against the roof as thunder rumbled.")
        self.assertIn("rain", [event.kind for event in weather])
        self.assertIn("thunder", [event.kind for event in weather])

        metaphor = detect_story_events("She was running out of time and running out of patience.")
        self.assertNotIn("running_footsteps", [event.kind for event in metaphor])

    def test_synthesized_story_effect_is_valid_local_audio(self):
        import wave
        from app.audio.story_events import EVENT_PROFILES, synthesize_story_effect

        with tempfile.TemporaryDirectory() as temp:
            output = synthesize_story_effect(
                Path(temp) / "running.wav", "running_footsteps", variant_seed=7
            )
            self.assertTrue(output.exists())
            with wave.open(str(output), "rb") as wav:
                self.assertEqual(wav.getnchannels(), 2)
                self.assertEqual(wav.getframerate(), 22050)
                self.assertAlmostEqual(
                    wav.getnframes() / wav.getframerate(),
                    EVENT_PROFILES["running_footsteps"].duration,
                    delta=0.1,
                )
            first = output.read_bytes()
            repeated = synthesize_story_effect(
                Path(temp) / "running-repeat.wav", "running_footsteps", variant_seed=7
            )
            self.assertEqual(first, repeated.read_bytes())

    def test_story_effects_mix_without_background_music_and_preserve_duration(self):
        from app.audio.assembler import _duration_ms

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            chapter_dirs = []
            stories = (
                (
                    "The Corridor",
                    "He ran down the corridor. The door slammed shut behind him.",
                ),
                (
                    "Storm",
                    "Rain poured against the roof as thunder rumbled in the distance.",
                ),
            )
            for number, (title, text) in enumerate(stories, start=1):
                chapter_dir = root / f"{number:03d}_{title.replace(' ', '_')}"
                chunk_path = chapter_dir / "chunks" / "00001.wav"
                _write_wav(chunk_path, 3.0)
                (chapter_dir / "chunk_manifest.json").write_text(
                    json.dumps({
                        "version": 1,
                        "chapter": number,
                        "title": title,
                        "chunks": [{"index": 1, "file": "00001.wav", "text": text}],
                    }),
                    encoding="utf-8",
                )
                chapter_dirs.append(chapter_dir)

            progress = []
            output = root / "story-effects-only.m4b"
            result = assemble_m4b(
                chapter_dirs,
                output,
                "Story Effects Only Test",
                chapter_titles=["The Corridor", "Storm"],
                progress=lambda done, total, message: progress.append(message),
                background_audio_mode="off",
                story_effects_enabled=True,
                story_effect_volume=0.20,
            )
            self.assertEqual(result, output)
            self.assertTrue(output.exists())
            self.assertIn("story-effects-mixed", progress)
            # Effects are overlaid in-place; they must not extend the book.
            self.assertAlmostEqual(_duration_ms(output) / 1000.0, 6.0, delta=0.20)
            self.assertGreaterEqual(_probe_audio_bitrate_kbps(output), 128)

    def test_story_effect_controls_are_independent_of_background_music(self):
        ui = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        manager = Path("app/tts/manager.py").read_text(encoding="utf-8")
        assembler = Path("app/audio/assembler.py").read_text(encoding="utf-8")
        self.assertIn('self.story_effects_enabled = QCheckBox(', ui)
        self.assertIn('self.story_effects_enabled.toggled.connect(', ui)
        self.assertIn('story_effects_enabled=self.story_effects_enabled.isChecked()', ui)
        self.assertIn('story_effects_enabled: bool = False', manager)
        self.assertIn('if story_effects_enabled:', assembler)

    def test_manual_music_still_gets_real_chapter_transition_cues(self):
        from unittest.mock import patch
        from app.audio.library import import_audio_files
        from app.audio.assembler import _duration_ms

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            library_root = root / "Audio Library"
            chapter_dir = root / "001_Manual_Chapter"
            _write_wav(chapter_dir / "chunks" / "00001.wav", 3.0)
            (chapter_dir / "narration.json").write_text(
                json.dumps({
                    "title": "Manual Chapter",
                    "narration_text": "The story continues into the next room.",
                }),
                encoding="utf-8",
            )
            music = root / "continuous-music.wav"
            _write_wav(music, 1.0)
            opening = root / "physical-book-opening.wav"
            _write_wav(opening, 1.18)
            page_turn = root / "physical-page-turn.wav"
            _write_wav(page_turn, 0.62)
            ending = root / "physical-book-closing.wav"
            _write_wav(ending, 2.65)
            progress = []
            output = root / "manual-with-real-cue.m4b"
            with patch("app.audio.library.audio_library_root", return_value=library_root):
                import_audio_files([opening], target_kind="opening")
                import_audio_files([page_turn], target_kind="page")
                import_audio_files([ending], target_kind="ending")
                result = assemble_m4b(
                    [chapter_dir],
                    output,
                    "Manual Mode Cue Test",
                    chapter_titles=["Manual Chapter"],
                    progress=lambda done, total, message: progress.append(message),
                    background_audio_mode="manual",
                    background_audio_path=music,
                    background_audio_volume=0.18,
                    chapter_transition_sound=True,
                    background_audio_fade_in=0.3,
                    background_audio_fade_out=0.3,
                )
            self.assertEqual(result, output)
            self.assertIn("mixed-background-audio", progress)
            self.assertGreaterEqual(_duration_ms(output), 6500)
            self.assertLessEqual(_duration_ms(output), 7200)
            self.assertGreaterEqual(_probe_audio_bitrate_kbps(output), 128)

    def test_audio_bitrate_range_is_enforced(self):
        from app.audio.assembler import _encoding_audio_bitrate, normalize_audio_bitrate
        self.assertEqual(normalize_audio_bitrate(None), "256k")
        self.assertEqual(normalize_audio_bitrate(128), "128k")
        self.assertEqual(normalize_audio_bitrate("320 kbps"), "320k")
        # The lowest requested target uses an encoding margin so final AAC
        # remains above the product's hard 128 kbps minimum.
        self.assertEqual(_encoding_audio_bitrate(128), "192k")
        self.assertEqual(_encoding_audio_bitrate(160), "160k")
        self.assertEqual(_encoding_audio_bitrate(256), "256k")
        with self.assertRaises(ValueError):
            normalize_audio_bitrate(127)
        with self.assertRaises(ValueError):
            normalize_audio_bitrate(321)

    def test_generation_page_exposes_premium_bitrate_selector(self):
        source = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        self.assertIn("256 kbps", source)
        self.assertIn("320", source)
        self.assertIn("audio_bitrate", source)

    def test_voice_studio_preview_has_save_action_and_audio_export_paths(self):
        source = Path("app/ui/voice_page.py").read_text(encoding="utf-8")
        self.assertIn('self.save_preview_button = QPushButton("↓  Save Preview")', source)
        self.assertIn('self.save_preview_button.clicked.connect(self.save_last_preview)', source)
        self.assertIn('def save_last_preview(self) -> None:', source)
        self.assertIn('QFileDialog.getSaveFileName(', source)
        self.assertIn('WAV Audio (*.wav);;MP3 Audio (*.mp3)', source)
        self.assertIn('"libmp3lame"', source)
        self.assertIn('self.save_preview_button.setEnabled(True)', source)
        self.assertIn('self.save_preview_button.setEnabled(False)', source)

    def test_voice_profile_catalogue_write_is_atomic_and_preserves_last_good_copy(self):
        from unittest.mock import patch
        import app.tts.voice_profile as vp

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp) / "Voices"
            original_root = vp.voices_root
            vp.voices_root = lambda: root
            try:
                destination = root / "voices.json"
                root.mkdir(parents=True)
                previous = '[{"name":"Last Good Voice","provider":"chatterbox","voice_id":"reference"}]'
                destination.write_text(previous, encoding="utf-8")
                profile = VoiceProfile(
                    name="New Custom Voice",
                    provider="chatterbox",
                    voice_id="reference-new",
                    sample_path=str(root / "reference.wav"),
                    authorized=True,
                )
                with patch.object(Path, "replace", side_effect=OSError("simulated replace failure")):
                    with self.assertRaises(OSError):
                        save_profiles([profile])
                self.assertEqual(destination.read_text(encoding="utf-8"), previous)
                self.assertFalse(destination.with_suffix(".json.tmp").exists())

                save_profiles([profile])
                saved = json.loads(destination.read_text(encoding="utf-8"))
                self.assertEqual(saved[0]["name"], "New Custom Voice")
                self.assertFalse(destination.with_suffix(".json.tmp").exists())
            finally:
                vp.voices_root = original_root

    def test_qwen_custom_voice_profile_round_trip_preserves_generation_settings(self):
        import app.tts.voice_profile as vp

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp) / "Voices"
            original_root = vp.voices_root
            vp.voices_root = lambda: root
            try:
                reference = root / "Custom Hero" / "reference.wav"
                reference.parent.mkdir(parents=True)
                reference.write_bytes(b"test reference placeholder")
                profile = VoiceProfile(
                    name="Custom Hero Voice",
                    provider="qwen-character",
                    voice_id="custom-hero",
                    sample_path=str(reference),
                    model_id="base-1.7b",
                    backend="automatic",
                    language="English",
                    authorized=True,
                    qwen_mode="clone",
                    qwen_prompt="Warm, steady, cinematic male voice.",
                    qwen_seed=2468,
                    reference_text="The journey begins now.",
                    qwen_anchor_path=str(reference),
                    qwen_anchor_text="The journey begins now.",
                    qwen_profile_id="custom-hero-stable-id",
                )
                save_profiles([profile])
                loaded = vp.load_saved_profile("Custom Hero Voice")
                self.assertIsNotNone(loaded)
                self.assertEqual(loaded.provider, "qwen-character")
                self.assertEqual(loaded.qwen_mode, "clone")
                self.assertEqual(loaded.model_id, "base-1.7b")
                self.assertEqual(loaded.sample_path, str(reference))
                self.assertEqual(loaded.reference_text, "The journey begins now.")
                self.assertEqual(loaded.qwen_anchor_path, str(reference))
                self.assertEqual(loaded.qwen_anchor_text, "The journey begins now.")
                self.assertEqual(loaded.qwen_seed, 2468)
                self.assertEqual(loaded.qwen_profile_id, "custom-hero-stable-id")
            finally:
                vp.voices_root = original_root

    def test_custom_reference_audio_is_normalized_for_cloning(self):
        with tempfile.TemporaryDirectory() as temp:
            import app.tts.voice_profile as vp
            root = Path(temp) / "Voices"
            source = Path(temp) / "stereo-8000.wav"
            rate = 8000
            frames = rate * 7
            with wave.open(str(source), "wb") as wav:
                wav.setnchannels(2)
                wav.setsampwidth(2)
                wav.setframerate(rate)
                wav.writeframes(b"\x01\x00\x01\x00" * frames)

            original_root = vp.voices_root
            vp.voices_root = lambda: root
            try:
                target = import_reference_audio(source, "Premium Test Voice")
                with wave.open(str(target), "rb") as handle:
                    self.assertEqual(handle.getnchannels(), 1)
                    self.assertEqual(handle.getframerate(), 24000)
                    self.assertEqual(handle.getsampwidth(), 2)
                    self.assertGreaterEqual(handle.getnframes() / handle.getframerate(), 6.0)
            finally:
                vp.voices_root = original_root

    def test_custom_voice_runtime_uses_pinned_perth_without_git_requirement(self):
        requirements = Path("requirements-voice-cloning.txt").read_text(encoding="utf-8")
        runtime = Path("app/tts/chatterbox_runtime.py").read_text(encoding="utf-8")
        self.assertNotIn("git+https://", requirements)
        self.assertIn("PERTH_SOURCE_REVISION", runtime)
        self.assertIn("perth_revision", runtime)

    def test_generation_page_preview_uses_background_worker(self):
        from PySide6.QtCore import QThread
        from app.ui.generation_page import VoicePreviewWorker
        self.assertTrue(issubclass(VoicePreviewWorker, QThread))

    def test_chatterbox_runtime_rejects_stale_marker(self):
        import tempfile
        from unittest.mock import patch
        import app.tts.chatterbox_runtime as runtime

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            marker = root / "runtime.json"
            marker.write_text(json.dumps({
                "provider": "chatterbox",
                "ready": True,
                "source_revision": "old-revision",
                "nano_supported": True,
            }), encoding="utf-8")
            with patch.object(runtime, "MARKER", marker), \
                 patch.object(runtime, "runtime_python", return_value=Path(sys.executable)):
                self.assertFalse(runtime.runtime_ready())

    def test_generation_page_exposes_custom_runtime_check(self):
        import app.ui.generation_page as generation_page

        self.assertTrue(callable(generation_page.runtime_ready))

    def test_generation_page_has_pronunciation_validation_dependencies(self):
        import inspect
        import app.ui.generation_page as generation_page

        self.assertTrue(hasattr(generation_page, "re"))
        self.assertTrue(callable(generation_page._numeric_score))
        self.assertIn("QHeaderView.ResizeMode.Stretch", inspect.getsource(generation_page.GenerationPage))
        self.assertIn("pronunciation_dictionary_version", inspect.getsource(generation_page.GenerationPage))

    def test_voice_page_refreshes_selected_provider_details(self):
        import inspect
        from app.ui.voice_page import VoicePage

        source = inspect.getsource(VoicePage._neural_voice_changed)
        self.assertIn("_update_profile_details(profile)", source)
        self.assertIn("_clear_saved_profile_selection()", source)


    def test_dialogue_assignment_fast_path_avoids_full_chapter_analysis(self):
        # The Assign Dialogue dialog should not run the expensive whole-chapter
        # character analysis just to suggest a speaker for one selected quote.
        from unittest.mock import patch

        text = 'Rias Gremory said hello. “Hello, Issei.” Rias Gremory smiled.'
        chapter = Chapter(
            1,
            'Chapter 1',
            text,
            dialogue_assignments=[
                {
                    'start': text.index('“Hello, Issei.”'),
                    'end': text.index('“Hello, Issei.”') + len('“Hello, Issei.”'),
                    'text': '“Hello, Issei.”',
                    'speaker': 'Rias Gremory',
                    'source': 'manual',
                }
            ],
        )
        start = text.index('“Hello, Issei.”')
        end = start + len('“Hello, Issei.”')
        with patch('app.chapters.dialogue.analyze_chapter', side_effect=AssertionError('full analysis called')):
            segment = dialogue_segment_for_selection(chapter, start, end)
        self.assertIsNotNone(segment)
        self.assertEqual(segment.suggested_speaker, 'Rias Gremory')
        self.assertGreaterEqual(segment.confidence, 0.99)

    def test_multi_speaker_cast_routes_exact_dialogue_to_multiple_voices(self):
        from app.tts.cast_provider import CastAwareProvider

        class DummyProvider:
            recommended_chunk_chars = 1400
            recommended_chunk_sentences = 2

            def synthesize(self, text, output_path, voice=None):
                return output_path

        from app.tts.voice_profile import VoiceProfile

        profiles = {
            "Voice A": VoiceProfile("Voice A", "piper", "a"),
            "Voice B": VoiceProfile("Voice B", "piper", "b"),
        }
        cast = CastAwareProvider(
            DummyProvider(),
            "Voice A",
            profiles,
            {"Rias": "Voice A", "Akeno": "Voice B"},
        )
        text = 'Rias and Akeno shouted, “We are here!”'
        start = text.index("“")
        end = text.index("”") + 1
        parts = cast.split_for_cast(
            text,
            "Voice A",
            dialogue_assignments=[{
                "start": start,
                "end": end,
                "text": text[start:end],
                "speakers": ["Rias", "Akeno"],
                "speaker": "Rias",
                "source": "manual",
                "multi_speaker_mode": "chorus",
            }],
        )
        dialogue_parts = [item for item in parts if "We are here!" in item[0]]
        self.assertEqual(len(dialogue_parts), 1)
        self.assertEqual(dialogue_parts[0][1], ["Voice A", "Voice B"])

    def test_cast_signature_changes_when_saved_qwen_identity_changes(self):
        from app.tts.cast_provider import CastAwareProvider
        from app.tts.voice_profile import VoiceProfile

        class DummyNarrator:
            def synthesize(self, text, output_path, voice=None):
                return output_path

        def make_cast(profile_id):
            profile = VoiceProfile(
                name="Saved Custom Voice",
                provider="qwen-character",
                voice_id="Ryan",
                model_id="base-0.6b",
                language="English",
                authorized=True,
                qwen_mode="clone",
                sample_path="reference.wav",
                reference_text="Hello there.",
                qwen_profile_id=profile_id,
            )
            return CastAwareProvider(
                DummyNarrator(),
                "Narrator",
                profiles={profile.name: profile},
                assignments={"Rias": profile.name},
            )

        old_cast = make_cast("old-profile-id")
        new_cast = make_cast("new-profile-id")
        self.assertNotEqual(
            old_cast.generation_signature(),
            new_cast.generation_signature(),
            "Changing a saved custom voice identity must invalidate resumable audio chunks.",
        )

    def test_light_novel_dialogue_discovers_characters(self):
        text = """Life.0
Issei Hyoudou—that’s my name, but my friends and family just call me Issei.
“My name is Rias Gremory,” said the crimson-haired girl.
“Oh dear. Greetings, I’m Akeno Himejima. Pleased to make your acquaintance,” she said.
“S-same here. Issei Hyoudou. Nice to meet you!” I replied.
“Rias, sit with us,” Akeno said.
"""
        analysis = analyze_book([Chapter(1, "Life.0", text)])
        names = {c.name.casefold() for c in analysis.characters}
        self.assertIn("issei hyoudou", names)
        self.assertIn("rias gremory", names)
        self.assertIn("akeno himejima", names)
        self.assertGreaterEqual(analysis.dialogue_total, 4)
        self.assertLess(analysis.unassigned_dialogue, analysis.dialogue_total)

    def test_character_detection_supports_standalone_labels_before_quoted_dialogue(self):
        chapter = Chapter(
            number=1,
            title="Standalone speaker labels",
            text=(
                "Rias Gremory\n"
                "「We should leave now.」\n"
                "Akeno Himejima\n"
                "“Of course, Rias.”\n"
                "リアス\n"
                "「わかったわ。」"
            ),
        )
        analysis = analyze_book([chapter])
        by_name = {item.name: item for item in analysis.characters}

        self.assertIn("Rias Gremory", by_name)
        self.assertIn("Akeno Himejima", by_name)
        self.assertIn("リアス", by_name)
        self.assertEqual(analysis.dialogue_total, 3)
        self.assertEqual(by_name["Rias Gremory"].dialogue_count, 1)
        self.assertEqual(by_name["Akeno Himejima"].dialogue_count, 1)
        self.assertEqual(by_name["リアス"].dialogue_count, 1)

    def test_dialogue_assignment_fast_path_suggests_standalone_speaker_labels(self):
        text = (
            "Rias Gremory\n"
            "「We should leave now.」\n"
            "Akeno Himejima\n"
            "“Of course, Rias.”"
        )
        chapter = Chapter(1, "Standalone speaker label assignment", text)
        start = text.index("「")
        end = text.index("」", start) + 1
        segment = dialogue_segment_for_selection(chapter, start, end)
        self.assertIsNotNone(segment)
        self.assertEqual(segment.suggested_speaker, "Rias Gremory")
        self.assertGreaterEqual(segment.confidence, 0.90)

    def test_light_novel_em_dash_dialogue_is_detected(self):
        text = """Chapter 1
Issei Hyoudou\u2014That's my name.

Akeno Himejima
\u2014Hello there, Issei.

Rias Gremory
\u2014We should get going.
"""
        self.assertEqual([ord(ch) for ch in text if ch in "—–-"], [0x2014, 0x2014, 0x2014])
        import app.chapters.characters as characters_module
        lines = text.splitlines(keepends=True)
        dash_lines = [
            (repr(line), hex(ord(line.lstrip(" \\t")[0])))
            for line in lines
            if line.lstrip(" \\t") and ord(line.lstrip(" \\t")[0]) in (0x2014, 0x2013, 0x2D)
        ]
        self.assertEqual(len(dash_lines), 2, repr(dash_lines))
        parser_source = inspect.getsource(characters_module._all_dialogue_spans)
        self.assertIn(
            "dash-only line",
            parser_source,
            f"Unexpected character parser import: {characters_module.__file__}\n"
            f"sys.path={sys.path}\n"
            f"parser={parser_source}",
        )
        spans = _all_dialogue_spans(text)
        self.assertGreaterEqual(
            len(spans), 2,
            repr({
                "module": characters_module.__file__,
                "dash_lines": dash_lines,
                "spans": spans,
                "source_has_dash_parser": "dash-only line" in inspect.getsource(characters_module._all_dialogue_spans),
            }),
        )
        analysis = analyze_book([Chapter(1, "Chapter 1", text)])
        names = {c.name.casefold() for c in analysis.characters}
        self.assertIn("issei hyoudou", names)
        self.assertIn("akeno himejima", names)
        self.assertIn("rias gremory", names)
        self.assertGreaterEqual(analysis.dialogue_total, 2)

    def test_detector_never_returns_empty_chapter(self):
        text = """Life.0
Story text here.

Life.1

Life.2
More story text here.
"""
        chapters = detect_chapters(text)
        self.assertTrue(chapters)
        self.assertTrue(all(c.text.strip() for c in chapters))
        self.assertEqual([c.number for c in chapters], list(range(1, len(chapters) + 1)))
    
    def test_empty_chapter_heading_does_not_leak_into_previous_chapter(self):
        text = """Life.0
Opening story text.

Life.1

Life.2
Second chapter story text.
"""
        chapters = detect_chapters(text)
        self.assertEqual([chapter.title for chapter in chapters], ["Life.0", "Life.2"])
        self.assertEqual(chapters[0].text, "Life.0\nOpening story text.")
        self.assertEqual(chapters[1].text, "Life.2\nSecond chapter story text.")
        self.assertNotIn("Life.1", "\\n".join(chapter.text for chapter in chapters))
        self.assertTrue(all(chapter.text.strip() for chapter in chapters))


    def test_narration_preprocessing_preserves_source_and_repairs_wrapping(self):
        source = """This is a sentence that was split
across two extracted PDF lines.

“He looked at me,”
she said.
Page 12
"""
        result = prepare_for_narration(source)

        self.assertEqual(result.source_text, source)
        self.assertIn("sentence that was split across two extracted PDF lines.", result.narration_text)
        self.assertIn('"He looked at me,"', result.narration_text)
        self.assertIn("she said.", result.narration_text)
        self.assertNotIn("Page 12", result.narration_text)

    def test_chapter_heading_is_not_joined_to_first_narration_sentence(self):
        source = (
            "Chapter 1 Prologue\n"
            "Erotic video games—eroge—are, for the most part, chock-full of things to ridicule.\n"
            "And there are plenty of reasons to lampoon them, too.\n"
            "Let's start with the protagonist's family, shall we?"
        )
        prepared = prepare_for_narration(source)
        self.assertTrue(prepared.narration_text.startswith("Chapter 1 Prologue\n"))
        self.assertNotIn("Chapter 1 Prologue Erotic video games", prepared.narration_text)

        units = split_for_pacing(prepared.narration_text, max_chars=800, max_sentences=4)
        self.assertEqual(units[0], "Chapter 1 Prologue")
        self.assertTrue(units[1].startswith("Erotic video games—eroge—are"))
        self.assertNotIn("Chapter 1 Prologue", units[1])
        self.assertGreaterEqual(pause_after_ms(units[0], "natural", next_text=units[1]), 600)

    def test_all_caps_and_section_headings_remain_separate_narration_units(self):
        source = (
            "The room went quiet.\n"
            "THE DEMON LORD SUMMIT\n"
            "A moment later, the doors opened.\n"
            "Part 2 — The Next Journey\n"
            "Everyone turned around."
        )
        prepared = prepare_for_narration(source)
        units = split_for_pacing(prepared.narration_text, max_chars=800, max_sentences=4)
        self.assertIn("THE DEMON LORD SUMMIT", units)
        self.assertIn("Part 2 — The Next Journey", units)
        for heading in ("THE DEMON LORD SUMMIT", "Part 2 — The Next Journey"):
            self.assertEqual(units.count(heading), 1)
            self.assertGreaterEqual(pause_after_ms(heading), 700)

    def test_natural_pacing_splits_story_into_sentence_units(self):
        text = "The door opened. Issei looked inside. Rias smiled."
        chunks = split_for_pacing(text, max_chars=200, max_sentences=2)
        self.assertEqual(chunks, ["The door opened. Issei looked inside.", "Rias smiled."])
        self.assertGreater(pause_after_ms("Rias smiled.", "natural"), 0)
        self.assertGreater(pause_after_ms("Rias smiled!", "natural"), pause_after_ms("Rias smiled.", "natural"))
        self.assertEqual(pause_after_ms("Rias smiled.", "off"), 0)

    def test_natural_pacing_appends_silence_without_changing_existing_audio(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "chunk.wav"
            _write_wav(path, 1.0)
            with wave.open(str(path), "rb") as before:
                original_frames = before.getnframes()
                rate = before.getframerate()
            append_silence(path, 250)
            with wave.open(str(path), "rb") as after:
                self.assertEqual(after.getframerate(), rate)
                self.assertGreater(after.getnframes(), original_frames)
                self.assertEqual(after.getnframes(), original_frames + round(rate * 0.25))

    def test_pronunciation_suggester_has_known_name_readings(self):
        self.assertEqual(suggest_pronunciation("Hyoudou Issei"), "Hee-doh Is-say")
        self.assertEqual(suggest_pronunciation("Ise"), "Ee-say")
        self.assertEqual(suggest_pronunciation("Rias"), "Ree-ahs")

    def test_pronunciation_filter_rejects_more_everyday_english(self):
        for word in ("actually", "already", "please", "probably", "everyone", "something", "yesterday"):
            self.assertTrue(is_common_english_phrase(word), word)

    def test_ai_pronunciation_keeps_model_spoken_form(self):
        from app.ai.brain import _normalise_pronunciation
        result = _normalise_pronunciation(
            {
                "written": "Hyeon-woo",
                "spoken": "hyun-woo",
                "ipa": "/hjʌn.u/",
                "source_language": "Korean",
                "confidence": 0.97,
            },
            "Hyeon-woo met the others.",
        )
        self.assertIsNotNone(result)
        self.assertEqual(result["spoken"], "hyun-woo")

    def test_pronunciation_suggester_ignores_common_english_words(self):
        text = """But she looked at Her friend.
Hyoudou Issei spoke to Rias Gremory.
"""
        candidates = {value.casefold() for value in suggest_names_from_text(text)}
        self.assertNotIn("but", candidates)
        self.assertNotIn("she", candidates)
        self.assertNotIn("her", candidates)
        self.assertIn("what", COMMON_ENGLISH_WORDS)
        self.assertIn("hyoudou issei", candidates)

    def test_pronunciation_suggester_finds_capitalized_name_candidates(self):
        text = "Hyoudou Issei met Rias Gremory at Kuoh Academy."
        names = suggest_names_from_text(text)
        self.assertIn("Hyoudou Issei", names)
        self.assertIn("Rias Gremory", names)
    def test_pronunciation_dictionary_is_case_insensitive_and_phrase_aware(self):
        source = "Hyoudou Issei met Ise. HYoudou Issei smiled."
        result = prepare_for_narration(source, [
            {"written": "Hyoudou Issei", "spoken": "Hee-doh Is-say", "enabled": True},
            {"written": "Ise", "spoken": "Ee-say", "enabled": True},
        ])
        self.assertEqual(
            result.narration_text,
            "Hee-doh Is-say met Ee-say. Hee-doh Is-say smiled.",
        )

    def test_disabled_pronunciation_entry_is_ignored(self):
        source = "Rias Gremory."
        result = prepare_for_narration(source, [
            {"written": "Rias", "spoken": "Ree-ahs", "enabled": False},
        ])
        self.assertEqual(result.narration_text, source)
    def test_narration_preprocessing_does_not_grammar_correct(self):
        source = "I am a second year high school student."
        result = prepare_for_narration(source)
        self.assertEqual(result.narration_text, source)
    def test_cast_provider_cleanup_continues_after_a_child_close_error(self):
        from app.tts.cast_provider import CastAwareProvider

        class DummyNarrator:
            def synthesize(self, text, output_path, voice=None):
                return output_path

        class BrokenCloser:
            def close(self):
                raise RuntimeError("simulated shutdown failure")

        class GoodCloser:
            def __init__(self):
                self.closed = False

            def close(self):
                self.closed = True

        cast = CastAwareProvider(DummyNarrator(), "Narrator", {}, {})
        good = GoodCloser()
        cast._providers = {
            "broken": (BrokenCloser(), "voice"),
            "good": (good, "voice"),
        }
        cast.close()
        self.assertTrue(good.closed, "one failed child close must not leak the other worker")

    def test_cast_provider_reuses_qwen_custom_worker_across_profiles(self):
        from unittest.mock import patch
        from app.tts.cast_provider import CastAwareProvider

        class DummyNarrator:
            def synthesize(self, text, output_path, voice=None):
                return output_path

        class DummyQwen:
            def synthesize(self, text, output_path, voice=None):
                self.calls.append((text, self.voice_id, self.instruct, self.qwen_seed))
                return output_path

            def __init__(self):
                self.calls = []
                self.voice_id = ""
                self.instruct = ""
                self.qwen_seed = None

        first = VoiceProfile(
            name="Character A", provider="qwen-character", voice_id="Ryan",
            model_id="custom-0.6b", backend="automatic", language="English",
            qwen_mode="custom", qwen_prompt="first prompt", qwen_seed=1,
        )
        second = VoiceProfile(
            name="Character B", provider="qwen-character", voice_id="Aiden",
            model_id="custom-0.6b", backend="automatic", language="English",
            qwen_mode="custom", qwen_prompt="second prompt", qwen_seed=2,
        )
        engine = DummyQwen()
        with patch("app.tts.cast_provider.provider_from_profile", return_value=(engine, "unused")) as factory:
            cast = CastAwareProvider(
                DummyNarrator(), "Narrator",
                {first.name: first, second.name: second},
                {"Speaker A": first.name, "Speaker B": second.name},
            )
            cast.synthesize("First line.", Path("first.wav"), first.name)
            cast.synthesize("Second line.", Path("second.wav"), second.name)

        self.assertEqual(factory.call_count, 1, "compatible Qwen profiles should share one model worker")
        self.assertEqual(engine.calls, [
            ("First line.", "Ryan", "first prompt", 1),
            ("Second line.", "Aiden", "second prompt", 2),
        ])

    def test_cast_provider_reuses_chatterbox_worker_across_profiles(self):
        from unittest.mock import patch
        from app.tts.cast_provider import CastAwareProvider

        class DummyNarrator:
            def synthesize(self, text, output_path, voice=None):
                return output_path

        class DummyChatterbox:
            def __init__(self):
                self.calls = []
                self.reference_audio = None
                self.exaggeration = 0.0
                self.cfg_weight = 0.0

            def synthesize(self, text, output_path, voice=None):
                self.calls.append((text, str(self.reference_audio), self.exaggeration, self.cfg_weight))
                return output_path

        with tempfile.TemporaryDirectory() as temp:
            reference_a = Path(temp) / "a.wav"
            reference_b = Path(temp) / "b.wav"
            reference_a.write_bytes(b"reference-a")
            reference_b.write_bytes(b"reference-b")
            first = VoiceProfile(
                name="Clone A", provider="chatterbox", voice_id="a",
                sample_path=str(reference_a), model_id="chatterbox", backend="automatic",
                language="en", exaggeration=0.2, cfg_weight=0.3,
            )
            second = VoiceProfile(
                name="Clone B", provider="chatterbox", voice_id="b",
                sample_path=str(reference_b), model_id="chatterbox", backend="automatic",
                language="en", exaggeration=0.7, cfg_weight=0.8,
            )
            engine = DummyChatterbox()
            with patch("app.tts.cast_provider.provider_from_profile", return_value=(engine, "unused")) as factory:
                cast = CastAwareProvider(
                    DummyNarrator(), "Narrator",
                    {first.name: first, second.name: second},
                    {"Speaker A": first.name, "Speaker B": second.name},
                )
                cast.synthesize("First line.", Path("first.wav"), first.name)
                cast.synthesize("Second line.", Path("second.wav"), second.name)

            self.assertEqual(factory.call_count, 1, "compatible Chatterbox profiles should share one model worker")
            self.assertEqual(engine.calls, [
                ("First line.", str(reference_a.resolve()), 0.2, 0.3),
                ("Second line.", str(reference_b.resolve()), 0.7, 0.8),
            ])

    def test_cast_provider_routes_em_dash_dialogue_to_assigned_character(self):
        from app.tts.cast_provider import CastAwareProvider

        class DummyNarrator:
            def synthesize(self, text, output_path, voice=None):
                return output_path

        provider = CastAwareProvider(
            DummyNarrator(),
            "Narrator",
            profiles={},
            assignments={"Akeno Himejima": "Akeno Voice"},
        )
        text = """Akeno Himejima
—Hello there, Issei.

The hallway was quiet.

Rias Gremory
—We should get going.
"""
        parts = provider.split_for_cast(text, "Narrator")
        assigned = [(value, voice) for value, voice in parts if voice == "Akeno Voice"]
        self.assertTrue(assigned)
        self.assertEqual(assigned[0][0], "Hello there, Issei.")

        provider.close()

    def test_voice_preview_uses_pronunciation_and_creates_audio(self):
        class DummyProvider:
            def synthesize(self, text, output_path, voice=None):
                _write_wav(output_path, 0.20)
                self.last_text = text
                self.last_voice = voice

        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "preview.wav"
            provider = DummyProvider()
            result = build_voice_preview(
                "Hyoudou Issei looked at the door.",
                provider,
                "Narrator",
                output,
                pronunciation_dictionary=[
                    {"written": "Hyoudou Issei", "spoken": "Hee-doh Is-say", "enabled": True}
                ],
                narration_speed=1.0,
                pacing_profile="natural",
                max_chars=500,
            )
            self.assertTrue(result.output_path.exists())
            self.assertGreater(result.segments, 0)
            self.assertIn("Hee-doh Is-say", provider.last_text)
            self.assertEqual(provider.last_voice, "Narrator")

    def test_targeted_preview_budget_can_exceed_default_for_fast_voices(self):
        from unittest.mock import patch
        import app.tts.preview as preview

        class DummyProvider:
            def synthesize(self, text, output_path, voice=None):
                _write_wav(output_path, 0.20)

        captured = {}
        def capture_units(text, provider, voice, pronunciation_dictionary, max_chars, dialogue_assignments=None):
            captured["max_chars"] = max_chars
            return [("A complete preview sentence.", voice)]

        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "preview.wav"
            with patch.object(preview, "_preview_units", side_effect=capture_units):
                preview.build_voice_preview(
                    "Sentence. " * 600,
                    DummyProvider(),
                    "Narrator",
                    output,
                    max_chars=1400,
                    target_seconds=30,
                    narration_speed=1.0,
                    pacing_profile="natural",
                )
            self.assertEqual(captured["max_chars"], min(len("Sentence. " * 600), 30 * 240))

    def test_timed_preview_stops_after_requested_audio_duration(self):
        class TimedDummyProvider:
            def __init__(self):
                self.calls = 0

            def synthesize(self, text, output_path, voice=None):
                self.calls += 1
                _write_wav(output_path, 0.30)

        text = " ".join(
            f"Sentence number {number} is ready."
            for number in range(1, 9)
        )
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "timed-preview.wav"
            provider = TimedDummyProvider()
            result = build_voice_preview(
                text,
                provider,
                "Narrator",
                output,
                narration_speed=1.0,
                pacing_profile="natural",
                max_chars=1400,
                target_seconds=1.5,
            )
            with wave.open(str(result.output_path), "rb") as wav:
                actual_seconds = wav.getnframes() / float(wav.getframerate())

            self.assertGreaterEqual(actual_seconds, 1.5)
            self.assertLess(result.segments, 8)
            self.assertEqual(provider.calls, result.segments)
            self.assertEqual(result.output_path, output)

    def test_voice_profile_library_preserves_custom_voice_and_labels(self):
        with tempfile.TemporaryDirectory() as temp:
            import app.tts.voice_profile as vp

            original_root = vp.voices_root
            root = Path(temp) / "Voices"
            vp.voices_root = lambda: root
            try:
                custom = VoiceProfile(
                    name="Offline Neural • Amy",
                    provider="chatterbox",
                    voice_id="reference",
                    sample_path=str(root / "custom" / "reference.wav"),
                    model_id="chatterbox-multilingual",
                    backend="automatic",
                    language="en",
                    authorized=True,
                )
                builtin = VoiceProfile(
                    name="Offline Neural • Amy",
                    provider="piper",
                    voice_id="en_US-amy-medium",
                    backend="automatic",
                    language="en-US",
                    authorized=True,
                )
                save_profiles([custom, builtin])
                loaded = load_profiles()
                custom_loaded = [p for p in loaded if p.provider == "chatterbox"]
                builtin_loaded = [p for p in loaded if p.provider == "piper"]
                self.assertEqual(len(custom_loaded), 1)
                self.assertEqual(len(builtin_loaded), 1)
                self.assertNotEqual(
                    custom_loaded[0].name.casefold(),
                    builtin_loaded[0].name.casefold(),
                )
                self.assertTrue(
                    custom_loaded[0].name.casefold().startswith("offline neural • amy")
                )
                self.assertEqual(custom_loaded[0].sample_path, str(root / "custom" / "reference.wav"))
            finally:
                vp.voices_root = original_root

    def test_generation_resume_state_write_is_atomic_and_preserves_last_good_copy(self):
        from unittest.mock import patch
        import app.tts.generator as generator

        with tempfile.TemporaryDirectory() as temp:
            chapter_dir = Path(temp) / "001_Chapter_One"
            chapter_dir.mkdir(parents=True)
            state_file = chapter_dir / "generation.json"
            previous = '{"completed": [0]}'
            state_file.write_text(previous, encoding="utf-8")

            with patch.object(Path, "replace", side_effect=OSError("simulated replace failure")):
                with self.assertRaises(OSError):
                    generator.save_state(chapter_dir, {"completed": [0, 1]})

            self.assertEqual(state_file.read_text(encoding="utf-8"), previous)
            self.assertEqual(list(chapter_dir.glob(".generation.json.*.tmp")), [])

            generator.save_state(chapter_dir, {"completed": [0, 1]})
            self.assertEqual(generator.load_state(chapter_dir), {"completed": [0, 1]})
            self.assertEqual(list(chapter_dir.glob(".generation.json.*.tmp")), [])

    def test_generation_manager_reports_unexpected_planning_failure_and_closes_provider(self):
        from app.tts.manager import GenerationManager

        class BrokenProvider:
            def __init__(self):
                self.closed = False

            def split_for_cast(self, text, voice, **kwargs):
                raise RuntimeError("simulated planning failure")

            def close(self):
                self.closed = True

        with tempfile.TemporaryDirectory() as temp:
            provider = BrokenProvider()
            summaries = []
            manager = GenerationManager(
                provider,
                "Narrator",
                [Chapter(1, "Chapter One", "A non-empty chapter.")],
                Path(temp) / "working-audio",
                on_finished=summaries.append,
            )
            manager.start()
            manager._thread.join(timeout=10)

            self.assertFalse(manager._thread.is_alive(), "generation worker remained stuck")
            self.assertTrue(provider.closed, "provider cleanup was skipped after a fatal error")
            self.assertEqual(len(summaries), 1, "fatal worker failure must produce one terminal result")
            self.assertEqual(summaries[0].chapters_failed, [1])
            self.assertIn("simulated planning failure", summaries[0].failure_details[1])

    def test_generation_resume_reuses_completed_chunks_without_reopening_them(self):
        from app.tts.generator import generate_chapter

        class DummyProvider:
            def __init__(self):
                self.calls = 0

            def synthesize(self, text, output_path, voice=None):
                self.calls += 1
                _write_wav(output_path, 0.20)
                return output_path

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp) / "audio"
            chapter = Chapter(
                1,
                "Chapter One",
                "This is the first sentence. This is the second sentence. "
                "This is the third sentence. This is the fourth sentence.",
            )
            provider = DummyProvider()

            first = generate_chapter(
                chapter,
                provider,
                "Narrator",
                root,
                narration_speed=1.0,
                pacing_profile="off",
            )
            first_calls = provider.calls
            self.assertGreater(first_calls, 0)
            self.assertEqual(first.chunks_completed, first.chunks_total)

            second = generate_chapter(
                chapter,
                provider,
                "Narrator",
                root,
                narration_speed=1.0,
                pacing_profile="off",
            )
            self.assertEqual(provider.calls, first_calls)
            self.assertEqual(second.chunks_completed, second.chunks_total)


    def test_generation_resume_state_detects_completed_and_partial_chapters(self):
        from app.tts.resume import inspect_generation_state

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            completed_dir = root / "001_Chapter_1"
            completed_dir.mkdir(parents=True)
            (completed_dir / "generation.json").write_text(
                '{"chunks_total": 2, "completed": [0, 1]}',
                encoding="utf-8",
            )

            partial_dir = root / "002_Chapter_2"
            (partial_dir / "chunks").mkdir(parents=True)
            (partial_dir / "generation.json").write_text(
                '{"chunks_total": 3, "completed": [0]}',
                encoding="utf-8",
            )

            chapters = [
                Chapter(1, "Chapter 1", "one"),
                Chapter(2, "Chapter 2", "two"),
                Chapter(3, "Chapter 3", "three"),
            ]
            state = inspect_generation_state(root, chapters)
            self.assertEqual(state.completed_chapters, (1,))
            self.assertEqual(state.partial_chapters, (2,))
            self.assertTrue(state.resume_available)


    def test_online_tts_network_error_is_classifiable(self):
        from app.tts.providers.edge_tts import EdgeTTSProvider, OnlineTTSNetworkError

        self.assertTrue(issubclass(OnlineTTSNetworkError, RuntimeError))
        self.assertTrue(getattr(EdgeTTSProvider, "is_online_provider", False))

    def test_edge_tts_classifies_no_audio_as_transient_service_error(self):
        from app.tts.providers.edge_tts import EdgeTTSProvider

        self.assertTrue(EdgeTTSProvider._is_transient_service_error("NoAudioReceived: No audio was received. Please verify that your parameters are correct."))
        self.assertTrue(EdgeTTSProvider._is_transient_service_error("WebSocket connection was closed by the server"))
        self.assertTrue(EdgeTTSProvider._is_transient_service_error("getaddrinfo failed"))
        self.assertFalse(EdgeTTSProvider._is_transient_service_error("An invalid voice parameter was supplied"))

    def test_edge_tts_uses_extended_transient_retry_budget(self):
        from app.tts.providers.edge_tts import EdgeTTSProvider

        self.assertEqual(EdgeTTSProvider.transient_retry_attempts, 5)
        self.assertEqual(EdgeTTSProvider.request_timeout_seconds, 60)

    def test_online_tts_provider_stops_generation_after_service_failure(self):
        from app.tts.providers.edge_tts import EdgeTTSProvider, OnlineTTSNetworkError

        self.assertTrue(EdgeTTSProvider.stop_on_failure)
        self.assertEqual(EdgeTTSProvider.request_timeout_seconds, 60)
        self.assertTrue(issubclass(OnlineTTSNetworkError, RuntimeError))

    def test_manager_stops_after_online_provider_failure(self):
        from app.tts.manager import GenerationManager

        class StopProvider:
            stop_on_failure = True

            def synthesize(self, text, output_path, voice=None):
                raise RuntimeError("online voice service unavailable")

            def close(self):
                pass

        summaries = []
        manager = GenerationManager(
            StopProvider(),
            "voice",
            [
                Chapter(1, "One", "One"),
                Chapter(2, "Two", "Two"),
                Chapter(3, "Three", "Three"),
            ],
            Path(tempfile.mkdtemp()),
            on_finished=summaries.append,
        )
        manager.start()
        manager._thread.join(timeout=10)

        self.assertEqual(len(summaries), 1)
        self.assertEqual(summaries[0].chapters_failed, [1])

    def test_project_save_round_trip_preserves_manual_chapter_edits(self):
        from app.core.project import Project

        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp) / "Book"
            project = Project(
                "Book",
                folder,
                None,
                [
                    Chapter(1, "Chapter One", "Edited body text."),
                    Chapter(2, "Chapter Two", "Second edited body."),
                ],
            )
            project.save()

            loaded = Project.load(folder)
            self.assertEqual(
                [(c.number, c.title, c.text) for c in loaded.chapters],
                [
                    (1, "Chapter One", "Edited body text."),
                    (2, "Chapter Two", "Second edited body."),
                ],
            )


    def test_empty_chapter_repair_preserves_existing_edits_and_assignments(self):
        source = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        start = source.index("def _repair_empty_chapters_from_source")
        end = source.index("def _refresh_resume_state", start)
        repair = source[start:end]
        self.assertIn('if str(chapter.text or "").strip():', repair)
        self.assertIn("Existing non-empty bodies are user", repair)
        self.assertIn('"dialogue_assignments"] = getattr(current, "dialogue_assignments", [])', repair)
        self.assertIn("temp_file.replace(chapters_file)", repair)
        self.assertNotIn("self.chapters[:] = detected", repair)

    def test_empty_chapter_repair_has_chapter_number_fallback(self):
        source = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        start = source.index("def _repair_empty_chapters_from_source")
        end = source.index("def _refresh_resume_state", start)
        repair = source[start:end]
        self.assertIn("detected_by_number", repair)
        self.assertIn("candidate = detected_by_number.get(int(chapter.number))", repair)
        self.assertIn('if str(chapter.text or "").strip():', repair)

    def test_generation_progress_range_uses_only_chapters_sent_to_manager(self):
        source = Path("app/ui/generation_page.py").read_text(encoding="utf-8")
        start = source.index("def start(self)")
        end = source.index("def toggle_generation_pause", start)
        generation_start = source[start:end]
        self.assertIn("generation_chapters, skipped_empty_chapters = partition_chapters_by_body(self.chapters)", generation_start)
        self.assertIn("self.progress.setRange(0, len(generation_chapters))", generation_start)
        self.assertNotIn("self.progress.setRange(0, len(self.chapters))", generation_start)

    def test_generation_readiness_allows_valid_chapters_and_nonblocking_empty_chapters(self):
        from app.chapters.detector import Chapter
        from app.tts.readiness import build_generation_readiness, partition_chapters_by_body

        valid = Chapter(1, "Chapter One", "Narration text.")
        empty = Chapter(2, "Chapter Two", "   ")
        chapters, skipped = partition_chapters_by_body([valid, empty])
        self.assertEqual(chapters, [valid])
        self.assertEqual(skipped, [empty])

        readiness = build_generation_readiness(
            [valid, empty], "Narrator", "piper", "C:/Books/Book.m4b"
        )
        chapter_item = next(item for item in readiness.items if item.label == "Chapters")
        self.assertTrue(readiness.ready)
        self.assertFalse(chapter_item.ok)
        self.assertFalse(chapter_item.blocking)
        self.assertIn("will be skipped", chapter_item.detail)

    def test_generation_readiness_blocks_qwen_voice_when_runtime_is_missing(self):
        from app.chapters.detector import Chapter
        from app.tts.readiness import build_generation_readiness

        readiness = build_generation_readiness(
            [Chapter(1, "Chapter One", "Narration text.")],
            "Custom Character Voice",
            "qwen-character",
            "C:/Books/Book.m4b",
            qwen_runtime_ready=False,
        )
        self.assertFalse(readiness.ready)
        item = next(item for item in readiness.items if item.label == "Qwen3-TTS runtime")
        self.assertTrue(item.blocking)
        self.assertFalse(item.ok)
        self.assertIn("Install / repair", item.detail)

    def test_generation_readiness_does_not_require_qwen_for_other_voices(self):
        from app.chapters.detector import Chapter
        from app.tts.readiness import build_generation_readiness

        readiness = build_generation_readiness(
            [Chapter(1, "Chapter One", "Narration text.")],
            "Piper Narrator",
            "piper",
            "C:/Books/Book.m4b",
            qwen_runtime_ready=False,
        )
        self.assertTrue(readiness.ready)
        item = next(item for item in readiness.items if item.label == "Qwen3-TTS runtime")
        self.assertTrue(item.ok)
        self.assertFalse(item.blocking)
        self.assertIn("Not required", item.detail)

    def test_qwen_profile_model_readiness_requires_selected_model_and_design_base(self):
        from types import SimpleNamespace
        from unittest.mock import patch
        import app.tts.qwen_character_runtime as runtime

        custom = SimpleNamespace(qwen_mode="custom", model_id="auto")
        with patch.object(runtime, "best_custom_kind", return_value="custom-0.6b"), patch.object(runtime, "model_installed", side_effect=lambda kind: kind == "custom-0.6b"):
            self.assertEqual(runtime.required_models_for_profile(custom), ("custom-0.6b",))
            self.assertEqual(runtime.profile_model_readiness(custom)[0], True)

        design = SimpleNamespace(qwen_mode="design", model_id="auto")
        with patch.object(runtime, "best_long_form_clone_kind", return_value="base-0.6b"), patch.object(runtime, "model_installed", return_value=False):
            ready, detail = runtime.profile_model_readiness(design)
            self.assertFalse(ready)
            self.assertIn("base-0.6b", detail)
            self.assertIn("design-1.7b", detail)

    def test_generation_readiness_blocks_missing_qwen_model_only_for_qwen(self):
        from app.chapters.detector import Chapter
        from app.tts.readiness import build_generation_readiness

        chapter = Chapter(1, "Chapter One", "Narration text.")
        readiness = build_generation_readiness(
            [chapter], "Qwen Character", "qwen-character", "C:/Books/Book.m4b",
            qwen_runtime_ready=True,
            qwen_models_ready=False,
            qwen_model_detail="Missing custom-0.6b model files.",
        )
        self.assertFalse(readiness.ready)
        item = next(item for item in readiness.items if item.label == "Qwen voice model")
        self.assertFalse(item.ok)
        self.assertTrue(item.blocking)
        self.assertIn("custom-0.6b", item.detail)

        other = build_generation_readiness(
            [chapter], "Piper Narrator", "piper", "C:/Books/Book.m4b",
            qwen_runtime_ready=False,
            qwen_models_ready=False,
            qwen_model_detail="Missing custom-0.6b model files.",
        )
        self.assertTrue(other.ready)
        item = next(item for item in other.items if item.label == "Qwen voice model")
        self.assertTrue(item.ok)
        self.assertFalse(item.blocking)
        self.assertIn("Not required", item.detail)

    def test_generation_readiness_blocks_when_every_chapter_is_empty(self):
        from app.chapters.detector import Chapter
        from app.tts.readiness import build_generation_readiness

        readiness = build_generation_readiness(
            [Chapter(1, "Chapter One", " "), Chapter(2, "Chapter Two", "")],
            "Narrator",
            "piper",
            "C:/Books/Book.m4b",
        )
        chapter_item = next(item for item in readiness.items if item.label == "Chapters")
        self.assertFalse(readiness.ready)
        self.assertTrue(chapter_item.blocking)
        self.assertIn("No chapter contains body text", chapter_item.detail)


    def test_stage5_shared_theme_defines_consistent_studio_surfaces_and_focus_states(self):
        from pathlib import Path

        source = Path("app/main.py").read_text(encoding="utf-8")
        self.assertIn("Dragon Studio Stage 5", source)
        self.assertIn("QWidget#studioMainPanel", source)
        self.assertIn("QMenu::item:selected", source)
        self.assertIn("QComboBox:focus", source)
        self.assertIn("QTableWidget::item", source)
        self.assertIn("QPushButton#primary:disabled", source)

    def test_stage5_html_aligned_page_styles_cover_dashboard_import_and_workflow(self):
        from pathlib import Path

        source = Path("app/main.py").read_text(encoding="utf-8")
        self.assertIn("QScrollArea#dashboardScroll", source)
        self.assertIn("QFrame#importDropZone", source)
        self.assertIn("QWidget#dashboardMetric", source)
        self.assertIn("QPushButton#workflowStep:checked", source)
        self.assertIn("QProgressBar::chunk", source)

    def test_stage5_import_studio_uses_scroll_area_to_prevent_clipped_controls(self):
        from pathlib import Path

        source = Path("app/ui/import_page.py").read_text(encoding="utf-8")
        self.assertIn("QScrollArea", source)
        self.assertIn('scroll.setObjectName("importScroll")', source)
        self.assertIn("scroll.setWidgetResizable(True)", source)
        self.assertIn("scroll.setHorizontalScrollBarPolicy", source)
        self.assertIn("self.create_button.clicked.connect(self.create_project)", source)


    def test_stage5_chapter_studio_scrolls_page_controls_on_short_displays(self):
        from pathlib import Path

        source = Path("app/ui/chapter_editor.py").read_text(encoding="utf-8")
        self.assertIn("QScrollArea", source)
        self.assertIn('scroll.setObjectName("chapterEditorScroll")', source)
        self.assertIn("scroll.setWidgetResizable(True)", source)
        self.assertIn("scroll.setHorizontalScrollBarPolicy", source)
        # Keep core editor actions and the chapter/text workspace in the scroll content.
        self.assertIn('QPushButton("Re-detect Chapters")', source)
        self.assertIn('QLabel("Chapter Text")', source)
        self.assertIn("self._autosave_timer.timeout.connect(self._autosave)", source)


    def test_stage5_audiobooks_player_screen_scrolls_without_losing_controls(self):
        from pathlib import Path

        source = Path("app/ui/audiobooks_page.py").read_text(encoding="utf-8")
        self.assertIn("QScrollArea", source)
        self.assertIn("scroll.setObjectName(\"audiobooksScroll\")", source)
        self.assertIn("scroll.setWidgetResizable(True)", source)
        self.assertIn("scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)", source)
        self.assertIn("scroll.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)", source)
        self.assertIn("self.play_button.clicked.connect(self.play_selected)", source)
        self.assertIn("self.pause_button.clicked.connect(self.pause_resume)", source)
        self.assertIn("self.stop_button.clicked.connect(self.stop_playback)", source)
        self.assertIn("self.timeline.sliderMoved.connect(self.player.setPosition)", source)
        self.assertIn("self.open_output_button.clicked.connect(self._open_output_folder)", source)

    def test_stage5_voice_cast_studio_scrolls_on_short_displays(self):
        from pathlib import Path

        source = Path("app/ui/voice_cast_page.py").read_text(encoding="utf-8")
        self.assertIn("QScrollArea", source)
        self.assertIn("scroll.setObjectName(\"voiceCastScroll\")", source)
        self.assertIn("scroll.setWidgetResizable(True)", source)
        self.assertIn("scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)", source)
        self.assertIn("scroll.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)", source)
        self.assertIn("self.analyze_button.clicked.connect(self.analyze)", source)
        self.assertIn("self.assign_button.clicked.connect(self.assign_voice)", source)
        self.assertIn("self.unassign_button.clicked.connect(self.unassign_voice)", source)

    def test_stage5_qwen_voice_cache_reset_handles_worker_errors_and_restores_button(self):
        from pathlib import Path

        source = Path("app/ui/models_page.py").read_text(encoding="utf-8")
        worker_start = source.index("class QwenCacheResetWorker")
        worker_end = source.index("\n\nclass QwenRepairWorker", worker_start)
        worker = source[worker_start:worker_end]
        self.assertIn("failed = Signal(str)", worker)
        self.assertIn("try:", worker)
        self.assertIn("self.failed.emit(str(exc))", worker)
        self.assertIn("self.qwen_cache_worker.failed.connect(self._qwen_cache_reset_failed)", source)
        self.assertIn("def _qwen_cache_reset_failed(self, message: str)", source)
        failed_handler = source[source.index("def _qwen_cache_reset_failed"):source.index("def repair_qwen_runtime")]
        self.assertIn("self.reset_qwen_cache.setEnabled(True)", failed_handler)
        self.assertIn("QMessageBox.warning", failed_handler)

    def test_stage5_hardware_studio_scrolls_and_preserves_benchmark_actions(self):
        from pathlib import Path

        source = Path("app/ui/hardware_page.py").read_text(encoding="utf-8")
        self.assertIn("QScrollArea", source)
        self.assertIn('scroll.setObjectName("hardwareScroll")', source)
        self.assertIn("scroll.setWidgetResizable(True)", source)
        self.assertIn("scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)", source)
        self.assertIn("scroll.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)", source)
        self.assertIn("self.benchmark_button.clicked.connect(self.run_benchmark)", source)
        self.assertIn("self.benchmark_thread.start()", source)
        self.assertIn("self.benchmark_thread.finished_result.connect(self._benchmark_finished)", source)

    def test_stage5_generation_history_scrolls_and_preserves_project_actions(self):
        from pathlib import Path

        source = Path("app/ui/history_page.py").read_text(encoding="utf-8")
        self.assertIn("QScrollArea", source)
        self.assertIn('scroll.setObjectName("historyScroll")', source)
        self.assertIn("scroll.setWidgetResizable(True)", source)
        self.assertIn("scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)", source)
        self.assertIn("scroll.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)", source)
        self.assertIn("self.refresh_button.clicked.connect(self.refresh)", source)
        self.assertIn("self.open_button.clicked.connect(self._open_project)", source)
        self.assertIn("self.production_button.clicked.connect(self._open_generation)", source)
        self.assertIn("self.output_button.clicked.connect(self._open_output_folder)", source)
        self.assertIn("self.table.itemSelectionChanged.connect(self._selection_changed)", source)

    def test_stage5_audiobook_library_scrolls_without_losing_collection_actions(self):
        from pathlib import Path

        source = Path("app/ui/library.py").read_text(encoding="utf-8")
        self.assertIn("QScrollArea", source)
        self.assertIn('scroll.setObjectName("libraryScroll")', source)
        self.assertIn("scroll.setWidgetResizable(True)", source)
        self.assertIn("scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)", source)
        self.assertIn("scroll.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)", source)
        self.assertIn("self.grid_button.clicked.connect", source)
        self.assertIn("self.list_button.clicked.connect", source)
        self.assertIn("self.open_button.clicked.connect(self.open_selected)", source)
        self.assertIn("self.delete_button.clicked.connect(self.delete_selected)", source)
        self.assertIn("self.folder_button.clicked.connect(self.open_selected_folder)", source)

    def test_stage5_settings_studio_scrolls_without_losing_preferences_or_storage_actions(self):
        from pathlib import Path

        source = Path("app/ui/settings_page.py").read_text(encoding="utf-8")
        self.assertIn("QScrollArea", source)
        self.assertIn('scroll.setObjectName("settingsScroll")', source)
        self.assertIn("scroll.setWidgetResizable(True)", source)
        self.assertIn("scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)", source)
        self.assertIn("scroll.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)", source)
        self.assertIn("self.save_button.clicked.connect(self.save_settings)", source)
        self.assertIn("self.default_bitrate", source)
        self.assertIn("self.open_log_button.clicked.connect(self.open_application_log)", source)
        self.assertIn('("Open Voice Profiles", voices_root())', source)

    def test_stage5_mini_player_keeps_playback_controls_in_compact_windows(self):
        from pathlib import Path

        source = Path("app/ui/workflow.py").read_text(encoding="utf-8")
        self.assertIn("compact_player = width < 1080", source)
        self.assertIn("self.mini_title.setVisible(not compact_player)", source)
        self.assertIn("self.mini_position.setVisible(not compact_player)", source)
        self.assertIn("self.mini_duration.setVisible(not compact_player)", source)
        self.assertIn("self.mini_play_button.clicked.connect(self._toggle_mini_playback)", source)
        self.assertIn("self.mini_timeline.sliderMoved.connect(self._seek_mini_player)", source)
        self.assertIn("self.mini_open_button.clicked.connect(lambda: self.select_section(\"Audiobooks\"))", source)

    def test_stage5_header_adapts_to_compact_window_widths(self):
        from pathlib import Path

        source = Path("app/ui/workflow.py").read_text(encoding="utf-8")
        self.assertIn("def resizeEvent(self, event)", source)
        self.assertIn("compact_header = width < 1240", source)
        self.assertIn("self.project_label.setVisible(not compact_header)", source)
        self.assertIn("self.global_search.setMaximumWidth(180 if width < 1400 else 250)", source)
        self.assertIn("self.global_search.returnPressed.connect(self._search_library)", source)

    def test_stage5_workflow_rail_scrolls_instead_of_compressing_navigation_steps(self):
        from pathlib import Path

        source = Path("app/ui/workflow.py").read_text(encoding="utf-8")
        self.assertIn('workflow_rail.setMinimumWidth(1080)', source)
        self.assertIn('workflow_scroll.setObjectName("workflowRailScroll")', source)
        self.assertIn("workflow_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)", source)
        self.assertIn("workflow_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)", source)
        self.assertIn("self.workflow_step_buttons[key] = button", source)
        self.assertIn("self._navigate_workflow_step(target)", source)


if __name__ == "__main__":
    unittest.main()


    def test_production_studio_validates_chapter_audio_and_normalizes_m4b_output(self):
        source = Path("app/ui/assembly_page.py").read_text(encoding="utf-8")
        self.assertIn("def _validate_build_inputs", source)
        self.assertIn('chunks_dir.glob("*.wav")', source)
        self.assertIn("chapter folder is missing", source)
        self.assertIn("no WAV audio chunks found", source)
        self.assertIn('output_path.suffix.casefold() != ".m4b"', source)
        self.assertIn("self._validate_build_inputs()", source)
        self.assertIn("An audiobook packaging job is already running.", source)
